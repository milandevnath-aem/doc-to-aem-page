"""Tests for ui/service.py -- the plain-Python pipeline wrapper behind the UI."""

from pathlib import Path

import docx

from ui.service import run_pipeline


def _build_docx(path: Path, builder) -> str:
    document = docx.Document()
    builder(document)
    document.save(str(path))
    return str(path)


def _valid_doc_builder(doc):
    doc.add_paragraph("[metadata]")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "Test Page"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "A test page"
    doc.add_paragraph("[component: hero-banner]")
    doc.add_paragraph("title:")
    doc.add_paragraph("Hero Title")


def test_run_pipeline_passes_clean_doc(tmp_path):
    path = _build_docx(tmp_path / "clean.docx", _valid_doc_builder)

    result = run_pipeline(path)

    assert result.status == "passed"
    assert result.filename == "clean.docx"
    assert result.issues == []
    assert result.ncm["blocks"][0]["name"] == "hero-banner"


def test_run_pipeline_blocks_parse_on_error_by_default(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "Only Title"
        # no "description" row -> missing-required-metadata error

    path = _build_docx(tmp_path / "broken.docx", build)

    result = run_pipeline(path)

    assert result.status == "failed"
    assert result.ncm is None
    assert any(i["code"] == "missing-required-metadata" for i in result.issues)


def test_run_pipeline_force_parses_despite_errors(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "Only Title"

    path = _build_docx(tmp_path / "broken.docx", build)

    result = run_pipeline(path, force=True)

    assert result.status == "failed"  # errors existed, status still reflects that
    assert result.ncm is not None  # but parsing was attempted anyway
    assert result.ncm["page"]["properties"]["title"] == "Only Title"


def test_run_pipeline_warnings_only_still_passes_and_parses(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "T"
        table.cell(2, 0).text, table.cell(2, 1).text = "description", "D"
        doc.add_paragraph("[component: hero-banner]")  # left empty -> "empty-component" warning
        doc.add_paragraph("[component: featurecards]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Real content")

    path = _build_docx(tmp_path / "warn_only.docx", build)

    result = run_pipeline(path)

    assert result.status == "passed"
    assert result.ncm is not None
    assert any(i["code"] == "empty-component" and i["severity"] == "warning" for i in result.issues)


def test_run_pipeline_unreadable_file_returns_failed_result(tmp_path):
    path = tmp_path / "not_really.docx"
    path.write_text("this is not a valid docx file")

    result = run_pipeline(str(path))

    assert result.status == "failed"
    assert result.ncm is None
    assert result.issues == [{
        "severity": "error",
        "code": "unreadable-file",
        "message": "'not_really.docx' could not be opened as a .docx file.",
        "paragraph_index": -1,
        "snippet": "",
        "context": "",
    }]
