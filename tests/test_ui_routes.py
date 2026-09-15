"""Tests for ui/app.py routes."""

import io
import re
import zipfile

from fastapi.testclient import TestClient

from ui import app as app_module
from ui import runs as runs_module
from ui.app import app
from ui.service import RunResult


def test_home_page_renders_upload_form():
    client = TestClient(app)

    response = client.get("/")

    assert response.status_code == 200
    assert 'name="docx_file"' in response.text
    assert 'action="/run"' in response.text
    # authoring-format selector: auto is the default, all three formats offered
    assert 'name="format"' in response.text
    assert 'value=""' in response.text          # auto
    assert 'value="v1v2"' in response.text
    assert 'value="v3"' in response.text


def test_home_page_shows_aem_site_config_section():
    client = TestClient(app)

    response = client.get("/")

    text = response.text
    # AEM config header + form + three action buttons
    assert "AEM site config" in text
    assert 'name="aem_base_url"' in text
    assert 'name="active_page_type"' in text
    assert 'name="auth_token_env"' in text
    assert 'action="/config/site"' in text
    assert 'action="/config/fetch"' in text
    assert 'action="/config/load-sample"' in text
    # default local AEM base URL is prefilled
    assert "http://localhost:4502" in text


def test_load_sample_route_writes_page_type_cache(tmp_path, monkeypatch):
    """POST /config/load-sample must materialize the bundled sample into
    the cache dir so subsequent pipeline runs read the same file."""
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    client = TestClient(app_module.app)
    response = client.post("/config/load-sample", follow_redirects=False)

    assert response.status_code == 303
    cached = cache_dir / "article.pagetypeconfig.json"
    assert cached.exists()
    text = cached.read_text(encoding="utf-8")
    assert '"template_path"' in text
    assert 'iffcotokio/components/text' in text  # from the bundled sample


def test_home_page_shows_page_and_container_resource_types(tmp_path, monkeypatch):
    """The two AEM-side resource types must be visible on the home page so
    operators can verify what AEM will resolve at install time.

    Uses an isolated cache dir so real fetched values on a dev box do not
    make the assertion brittle -- we force the bundled sample to be the
    active config for the test."""
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    client = TestClient(app_module.app)
    response = client.get("/")

    text = response.text
    assert "Page resource type" in text
    assert "Container resource type" in text
    # falls back to the bundled sample -- values from config/samples/article.pagetypeconfig.json
    assert "iffcotokio/components/page" in text
    assert "iffcotokio/components/parsys" in text
    assert "slug_from_url" in text


def test_save_site_config_preserves_other_page_types(tmp_path, monkeypatch):
    """Saving the form only touches the article CF path -- an already-saved
    blog/category entry must survive the round trip."""
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    site_path = cache_dir / "site.json"
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", site_path)

    # seed a site config with two page types already configured
    from aem_client.config import SiteConfig, save_site_config
    save_site_config(SiteConfig(
        aem_base_url="http://localhost:4502",
        active_page_type="article",
        auth_token_env="AEM_AUTH_TOKEN",
        cf_paths={
            "article": "/content/dam/x/article-old",
            "blog": "/content/dam/x/blog",   # must survive
        },
    ), site_path)

    client = TestClient(app_module.app)
    resp = client.post(
        "/config/site",
        data={
            "aem_base_url": "http://localhost:4502",
            "active_page_type": "article",
            "auth_token_env": "AEM_AUTH_TOKEN",
            "article_cf_path": "/content/dam/x/article-new",
        },
        follow_redirects=False,
    )
    assert resp.status_code == 303

    from aem_client.config import load_site_config
    reloaded = load_site_config(site_path)
    assert reloaded.cf_paths["article"] == "/content/dam/x/article-new"
    assert reloaded.cf_paths["blog"] == "/content/dam/x/blog"


def test_home_page_reflects_saved_site_config(tmp_path, monkeypatch):
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    client = TestClient(app_module.app)
    saved = client.post(
        "/config/site",
        data={
            "aem_base_url": "http://custom-aem:4503",
            "active_page_type": "article",
            "auth_token_env": "AEM_TOKEN_STAGE",
            "article_cf_path": "/content/dam/x/article",
        },
        follow_redirects=False,
    )
    assert saved.status_code == 303

    home = client.get("/")
    assert "http://custom-aem:4503" in home.text
    assert "AEM_TOKEN_STAGE" in home.text
    assert "/content/dam/x/article" in home.text


def test_run_route_forwards_format_override(tmp_path, monkeypatch):
    """Explicit format on the form must reach the parser and end up on the
    result page's detected-format badge."""
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)

    import docx as docx_lib
    document = docx_lib.Document()
    document.add_paragraph("[metadata]")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "T"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "D"
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    client = TestClient(app_module.app)
    response = client.post(
        "/run",
        files={"docx_file": ("sample.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"format": "v1v2"},
        follow_redirects=False,
    )

    assert response.status_code == 303
    result_response = client.get(response.headers["location"])
    assert "Detected format" in result_response.text
    assert "v1v2" in result_response.text


def test_run_route_uploads_and_redirects_to_result(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)

    import docx as docx_lib
    document = docx_lib.Document()
    document.add_paragraph("[metadata]")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "Test Page"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "A test page"
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    client = TestClient(app_module.app)
    response = client.post(
        "/run",
        files={"docx_file": ("sample.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    assert response.headers["location"].startswith("/runs/")


def test_full_upload_result_history_flow(tmp_path, monkeypatch):
    """End-to-end: upload a docx, follow the redirect to the result page,
    confirm parsed content renders, then confirm the run shows up on /history.
    """
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)

    import docx as docx_lib
    document = docx_lib.Document()
    document.add_paragraph("[metadata]")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "Test Page"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "A test page"
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    client = TestClient(app_module.app)
    response = client.post(
        "/run",
        files={"docx_file": ("e2e-flow.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    location = response.headers["location"]
    assert location.startswith("/runs/")

    result_response = client.get(location)

    assert result_response.status_code == 200
    assert "e2e-flow.docx" in result_response.text
    assert "Test Page" in result_response.text

    history_response = client.get("/history")

    assert history_response.status_code == 200
    assert "e2e-flow.docx" in history_response.text


def test_run_route_rejects_non_docx_file(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)
    client = TestClient(app_module.app)

    response = client.post(
        "/run",
        files={"docx_file": ("notes.txt", io.BytesIO(b"hello"), "text/plain")},
    )

    assert response.status_code == 400


def test_run_route_prevents_path_traversal(tmp_path, monkeypatch):
    """Verify that malicious filenames with path traversal segments are sanitized."""
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)

    import docx as docx_lib
    document = docx_lib.Document()
    document.add_paragraph("[metadata]")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "Test Page"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "A test page"
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    client = TestClient(app_module.app)
    # Upload with a crafted filename containing path traversal segments
    response = client.post(
        "/run",
        files={"docx_file": ("../../evil.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    # Verify that the file was written only inside tmp_path, not in parent directories
    uploaded_docx_files = list(tmp_path.glob("*.docx"))
    assert len(uploaded_docx_files) == 1, f"Expected 1 .docx file in UPLOADS_DIR, found {len(uploaded_docx_files)}"
    # Verify the file has the sanitized name (evil.docx, not ../../evil.docx)
    assert "evil.docx" in uploaded_docx_files[0].name
    # Verify no files leaked to parent directories
    parent = tmp_path.parent
    parent_files = list(parent.glob("evil.docx"))
    assert len(parent_files) == 0, f"Path traversal vulnerability: evil.docx found in {parent}"


def test_result_page_shows_status_and_issues():
    result = RunResult(
        status="failed",
        filename="broken.docx",
        issues=[{
            "severity": "error", "code": "missing-required-metadata",
            "message": "Metadata table is missing required field(s): ['description'].",
            "paragraph_index": 0, "snippet": "[metadata]", "context": "",
        }],
        ncm=None,
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}")

    assert response.status_code == 200
    assert "missing-required-metadata" in response.text
    assert "FAILED" in response.text


def test_result_page_renders_rte_html_preview():
    result = RunResult(
        status="passed",
        filename="clean.docx",
        issues=[],
        ncm={
            "page": {"properties": {"title": "T"}},
            "blocks": [{
                "name": "hero-banner", "variant": None,
                "fields": {"title": "Hero", "htmltext": "<p>Hello <strong>world</strong></p>"},
                "items": [], "table": None,
            }],
        },
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}")

    assert response.status_code == 200
    assert "<strong>world</strong>" in response.text  # rendered raw HTML, not escaped


def test_result_page_shows_full_ncm_json():
    result = RunResult(
        status="passed",
        filename="clean.docx",
        issues=[],
        ncm={
            "page": {"properties": {"title": "T"}},
            "blocks": [{
                "name": "hero-banner", "variant": None,
                "fields": {"title": "Hero", "description": "<p>Hello</p>"},
                "items": [], "table": None,
            }],
        },
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}")

    assert response.status_code == 200
    # The raw master JSON is dumped into a <pre> which HTML-autoescapes every
    # `"`, `<`, `>` -- so we assert on the HTML-source form. Once the browser
    # renders the page the reader sees the normal JSON with real `<p>` tags.
    text = response.text
    assert "&#34;name&#34;: &#34;hero-banner&#34;" in text
    assert "&lt;p&gt;Hello&lt;/p&gt;" in text
    # crucially: no more unicode escapes for HTML brackets (Jinja `tojson`
    # would have produced those; the `html_json` filter avoids them)
    assert "\\u003c" not in text


def test_result_page_renders_items_with_image_thumbnail():
    result = RunResult(
        status="passed",
        filename="v2.docx",
        issues=[],
        ncm={
            "page": {"properties": {"title": "T"}},
            "blocks": [{
                "name": "featurecards", "variant": None,
                "fields": {"title": "Benefits"},
                "items": [{
                    "id": None, "implicit": True,
                    "fields": {
                        "title": "Financial Protection",
                        "htmltext": "<p>Card body</p>",
                        "image": {"source": "embedded", "name": "image-1.png",
                                  "path": "ui/static/assets/run-x/image-1.png"},
                    },
                }],
                "table": None,
            }],
        },
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}")

    assert response.status_code == 200
    assert "Financial Protection" in response.text
    assert "<p>Card body</p>" in response.text
    assert '<img class="asset-thumb" src="/static/assets/run-x/image-1.png"' in response.text


def test_run_route_extracts_embedded_images_to_assets(tmp_path, monkeypatch):
    import base64

    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path / "uploads")
    monkeypatch.setattr(app_module, "ASSETS_DIR", tmp_path / "assets")
    (tmp_path / "uploads").mkdir()

    png_1px = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
    )
    img_file = tmp_path / "pic.png"
    img_file.write_bytes(png_1px)

    import docx as docx_lib
    document = docx_lib.Document()
    document.add_paragraph("[metadata]")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "Test Page"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "A test page"
    document.add_paragraph("[component: featurecards]")
    document.add_paragraph("Cards Title")
    document.add_picture(str(img_file))
    document.add_paragraph("Pasted Card")
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    client = TestClient(app_module.app)
    response = client.post(
        "/run",
        files={"docx_file": ("with-image.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )

    assert response.status_code == 303
    extracted = list((tmp_path / "assets").rglob("image-1.png"))
    assert len(extracted) == 1

    result_response = client.get(response.headers["location"])
    assert "Pasted Card" in result_response.text
    assert "image-1.png" in result_response.text


def test_result_page_shows_download_package_link_when_tree_present():
    from ui.service import RunResult
    result = RunResult(
        status="passed", filename="v3.docx", issues=[],
        ncm={"format": "v3", "page": {"properties": {"title": "T"}}, "blocks": []},
        aem_tree={"name": "t", "path": "/content/x/t",
                  "primary_type": "cq:Page", "properties": {}, "children": []},
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}")

    assert response.status_code == 200
    assert f'/runs/{run_id}/package.zip' in response.text
    assert "Download AEM Package" in response.text


def test_download_package_returns_installable_zip():
    from ui.service import RunResult
    result = RunResult(
        status="passed", filename="Sample Doc.docx", issues=[],
        ncm={"format": "v3", "page": {"properties": {"title": "Sample"}}, "blocks": []},
        aem_tree={
            "name": "sample", "path": "/content/site/en/articles/sample",
            "primary_type": "cq:Page", "properties": {},
            "children": [{
                "name": "jcr:content", "primary_type": "cq:PageContent",
                "properties": {"jcr:title": "Sample",
                               "sling:resourceType": "site/components/page"},
                "children": [],
            }],
        },
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}/package.zip")

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert "attachment" in response.headers["content-disposition"]
    # zip name derived from filename stem
    assert "sample-doc" in response.headers["content-disposition"].lower()

    import zipfile
    from io import BytesIO
    with zipfile.ZipFile(BytesIO(response.content)) as zf:
        names = set(zf.namelist())
        assert "META-INF/vault/filter.xml" in names
        assert "META-INF/vault/properties.xml" in names
        assert "jcr_root/content/site/en/articles/sample/.content.xml" in names
        page_xml = zf.read(
            "jcr_root/content/site/en/articles/sample/.content.xml"
        ).decode()
    assert 'jcr:primaryType="cq:Page"' in page_xml
    assert 'jcr:title="Sample"' in page_xml


def test_download_package_404_when_run_missing():
    client = TestClient(app_module.app)
    response = client.get("/runs/does-not-exist/package.zip")
    assert response.status_code == 404


def test_download_package_400_when_no_aem_tree():
    from ui.service import RunResult
    result = RunResult(
        status="passed", filename="v3.docx", issues=[],
        ncm={"format": "v3", "page": {"properties": {"title": "T"}}, "blocks": []},
        aem_tree=None,
    )
    run_id = runs_module.save_run(result)
    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}/package.zip")
    assert response.status_code == 400


def test_result_page_shows_aem_node_tree_when_page_type_config_active(tmp_path, monkeypatch):
    """End-to-end: upload a doc while the article sample config is active,
    result page must show the mapped AEM node tree with the target path.
    Cache dir is isolated so real fetched values on the dev box do not
    make the assertion brittle -- the pipeline falls back to the bundled
    sample."""
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path / "uploads")
    (tmp_path / "uploads").mkdir()

    import docx as docx_lib
    document = docx_lib.Document()
    document.add_paragraph().add_run("Cashless vs Reimbursement").bold = True
    document.add_paragraph("Meta Title: Cashless vs Reimbursement Guide")
    document.add_paragraph("Meta Description: A guide.")
    document.add_paragraph("New Required URL: https://www.iffcotokio.co.in/"
                           "health-insurance/cashless-vs-reimbursement")
    document.add_paragraph("Meta keywords: kw1, kw2")
    document.add_paragraph("Body content.")
    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)

    client = TestClient(app_module.app)
    posted = client.post(
        "/run",
        files={"docx_file": ("cashless.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        follow_redirects=False,
    )
    assert posted.status_code == 303

    page = client.get(posted.headers["location"])
    text = page.text
    assert "AEM Node Tree" in text
    # target JCR path = content_root + the URL's own section hierarchy + slug,
    # from the bundled sample (pageNameStrategy slug_from_url)
    assert "/content/iffcotokio/en/health-insurance/cashless-vs-reimbursement" in text
    # jcr:content properties visible in the JSON dump
    assert "iffcotokio/components/text" in text  # page-description mapped
    assert "cq:template" in text


def test_result_page_has_copy_buttons_per_section():
    result = RunResult(
        status="passed",
        filename="v3.docx",
        issues=[],
        ncm={
            "format": "v3",
            "page": {"properties": {"title": "T"}},
            "blocks": [{
                "name": "faq", "variant": None,
                "fields": {"title": "FAQS"},
                "items": [], "table": None,
            }],
        },
    )
    run_id = runs_module.save_run(result)

    client = TestClient(app_module.app)
    response = client.get(f"/runs/{run_id}")

    text = response.text
    import html as html_mod
    import json as json_mod
    import re

    # 3 copy buttons: Full NCM, Page Properties, one per block
    assert text.count('class="copy-btn"') == 3
    # each carries the section's JSON as data-copy. Autoescape encodes JSON
    # quotes as &#34; -- decode and JSON-parse to verify the payload survives
    payloads = [json_mod.loads(html_mod.unescape(m.group(1)))
                for m in re.finditer(r"data-copy='([^']*)'", text)]
    assert len(payloads) == 3
    # third payload = the faq block itself
    assert payloads[2]["name"] == "faq"
    assert payloads[2]["fields"]["title"] == "FAQS"
    assert 'navigator.clipboard' in text            # handler script present


def test_result_page_404_for_unknown_run():
    client = TestClient(app_module.app)

    response = client.get("/runs/does-not-exist")

    assert response.status_code == 404


def test_history_page_lists_past_runs():
    runs_module.save_run(RunResult(status="passed", filename="a.docx", issues=[], ncm=None))
    runs_module.save_run(RunResult(status="failed", filename="b.docx", issues=[{
        "severity": "error", "code": "x", "message": "m", "paragraph_index": 0, "snippet": "", "context": "",
    }], ncm=None))

    client = TestClient(app_module.app)
    response = client.get("/history")

    assert response.status_code == 200
    assert "a.docx" in response.text
    assert "b.docx" in response.text


def test_run_batch_route_uploads_multiple_docx_and_redirects(tmp_path, monkeypatch):
    """Client picks a folder with 2 subfolders, each holding one docx.
    Server receives files with webkitRelativePath filenames, groups them
    into one batch run, and redirects to /batch-runs/{id}.
    Cache dir is isolated so real fetched values on the dev box do not
    make this brittle -- the pipeline falls back to the bundled sample."""
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    import docx as docx_lib
    def make_v3(title, slug):
        d = docx_lib.Document()
        d.add_paragraph().add_run(title).bold = True
        d.add_paragraph(f"Meta Title: {title}")
        d.add_paragraph("Meta Description: desc")
        # the bundled sample runs slug_from_url -- no URL, no page
        d.add_paragraph("New Required URL: "
                        f"https://www.iffcotokio.co.in/health-insurance/{slug}")
        d.add_paragraph("Meta keywords: kw")
        d.add_paragraph("Body.")
        buf = io.BytesIO()
        d.save(buf)
        buf.seek(0)
        return buf

    files = [
        ("all_files", ("Articles/Article1/one.docx", make_v3("First", "first"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("all_files", ("Articles/Article2/two.docx", make_v3("Second", "second"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]

    client = TestClient(app_module.app)
    resp = client.post("/run-batch", files=files, follow_redirects=False)

    assert resp.status_code == 303
    assert resp.headers["location"].startswith("/batch-runs/")

    page = client.get(resp.headers["location"])
    assert "Batch Run" in page.text
    assert "Article1" in page.text
    assert "Article2" in page.text
    assert "one.docx" in page.text
    assert "two.docx" in page.text
    assert "PASSED" in page.text


def test_run_batch_route_rejects_when_no_docx(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)
    client = TestClient(app_module.app)
    resp = client.post(
        "/run-batch",
        files=[("all_files", ("readme.txt", io.BytesIO(b"hi"), "text/plain"))],
    )
    assert resp.status_code == 400


def test_download_batch_package_returns_multi_page_zip(tmp_path, monkeypatch):
    """End-to-end: POST 2 docx -> download /batch-runs/{id}/package.zip ->
    unzip -> both pages present + multi-root filter.xml.
    Cache dir is isolated so real fetched values on the dev box do not
    make this brittle -- the pipeline falls back to the bundled sample."""
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    import docx as docx_lib
    def make_v3(title, slug):
        d = docx_lib.Document()
        d.add_paragraph().add_run(title).bold = True
        d.add_paragraph(f"Meta Title: {title}")
        d.add_paragraph("Meta Description: desc")
        # the bundled sample runs slug_from_url -- no URL, no page
        d.add_paragraph("New Required URL: "
                        f"https://www.iffcotokio.co.in/health-insurance/{slug}")
        d.add_paragraph("Meta keywords: kw")
        d.add_paragraph("Body.")
        buf = io.BytesIO()
        d.save(buf)
        buf.seek(0)
        return buf

    files = [
        ("all_files", ("Articles/Article1/one.docx", make_v3("First Article", "first-article"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("all_files", ("Articles/Article2/two.docx", make_v3("Second Article", "second-article"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]
    client = TestClient(app_module.app)
    resp = client.post("/run-batch", files=files, follow_redirects=False)
    location = resp.headers["location"]
    run_id = location.rsplit("/", 1)[-1]

    zip_resp = client.get(f"/batch-runs/{run_id}/package.zip")
    assert zip_resp.status_code == 200
    assert zip_resp.headers["content-type"] == "application/zip"

    import zipfile
    from io import BytesIO
    with zipfile.ZipFile(BytesIO(zip_resp.content)) as zf:
        names = set(zf.namelist())
        filter_xml = zf.read("META-INF/vault/filter.xml").decode()

    # both pages present at their JCR paths, nested under the URL hierarchy
    assert ("jcr_root/content/iffcotokio/en/health-insurance/first-article/"
            ".content.xml") in names
    assert ("jcr_root/content/iffcotokio/en/health-insurance/second-article/"
            ".content.xml") in names
    # multi-root filter
    assert filter_xml.count("<filter root=") == 2


def test_batch_report_zip_downloads_both_csvs(tmp_path, monkeypatch):
    """POST 2 docx -> download /batch-runs/{id}/report.zip -> both CSVs
    present, header confirms column order. Cache dir isolated same as the
    batch package test above."""
    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    import docx as docx_lib
    def make_v3(title, slug):
        d = docx_lib.Document()
        d.add_paragraph().add_run(title).bold = True
        d.add_paragraph(f"Meta Title: {title}")
        d.add_paragraph("Meta Description: desc")
        # the bundled sample runs slug_from_url -- no URL, no page
        d.add_paragraph("New Required URL: "
                        f"https://www.iffcotokio.co.in/health-insurance/{slug}")
        d.add_paragraph("Meta keywords: kw")
        d.add_paragraph("Body.")
        buf = io.BytesIO()
        d.save(buf)
        buf.seek(0)
        return buf

    files = [
        ("all_files", ("Articles/Article1/one.docx", make_v3("First Article", "first-article"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
        ("all_files", ("Articles/Article2/two.docx", make_v3("Second Article", "second-article"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
    ]
    client = TestClient(app_module.app)
    resp = client.post("/run-batch", files=files, follow_redirects=False)
    location = resp.headers["location"]
    run_id = location.rsplit("/", 1)[-1]

    response = client.get(f"/batch-runs/{run_id}/report.zip")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    assert "attachment" in response.headers["content-disposition"]

    import zipfile
    from io import BytesIO
    with zipfile.ZipFile(BytesIO(response.content)) as zf:
        assert sorted(zf.namelist()) == ["faqs.csv", "pages.csv"]
        header = zf.read("pages.csv").decode("utf-8-sig").splitlines()[0]
        assert header.startswith("doc,folder,status,url,page_path")


def test_batch_report_zip_unknown_run_is_404():
    client = TestClient(app_module.app)
    assert client.get("/batch-runs/does-not-exist/report.zip").status_code == 404


def test_history_page_empty_state():
    client = TestClient(app_module.app)

    response = client.get("/history")

    assert response.status_code == 200
    assert "No runs yet" in response.text


def test_fetch_config_surfaces_aem_rejection_instead_of_silently_using_cache(
    tmp_path, monkeypatch,
):
    """A 401 from AEM must reach the operator. Before this, fetch_or_cached
    swallowed it, the route redirected to "/" as though it worked, and the
    pipeline kept running on a months-old cache -- which is how a batch of
    real documents shipped with an empty siteOrigin."""
    import httpx
    from aem_client.config import SiteConfig, save_site_config

    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")
    save_site_config(
        SiteConfig(aem_base_url="http://localhost:4502", active_page_type="article",
                   cf_paths={"article": "/content/dam/x/article"}),
        cache_dir / "site.json",
    )
    # seed a cache so the old fallback would have had something to hide behind
    monkeypatch.setattr(
        app_module, "_page_type_cache_path",
        lambda pt: cache_dir / f"{pt}.pagetypeconfig.json",
    )
    from aem_client.config import load_page_type_from_sample, save_page_type_config
    save_page_type_config(load_page_type_from_sample("article"),
                          cache_dir / "article.pagetypeconfig.json")

    def _deny(request):
        return httpx.Response(401)

    real_client = httpx.Client
    monkeypatch.setattr(
        app_module.httpx, "Client",
        lambda *a, **kw: real_client(transport=httpx.MockTransport(_deny)),
    )

    client = TestClient(app_module.app)
    response = client.post("/config/fetch", follow_redirects=False)

    assert response.status_code == 502
    detail = response.json()["detail"]
    assert "401" in detail
    assert "AEM_AUTH_TOKEN" in detail          # names the env var to set
    assert "unchanged" in detail               # says the cache was not damaged


def test_home_page_shows_when_the_cached_config_was_fetched(tmp_path, monkeypatch):
    """Without a date on screen there is nothing telling an operator the
    config they are looking at is older than their last AEM change."""
    from aem_client.config import load_page_type_from_sample, save_page_type_config

    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")
    monkeypatch.setattr(
        app_module, "_page_type_cache_path",
        lambda pt: cache_dir / f"{pt}.pagetypeconfig.json",
    )
    save_page_type_config(load_page_type_from_sample("article"),
                          cache_dir / "article.pagetypeconfig.json")

    text = TestClient(app_module.app).get("/").text
    assert "fetched 20" in text and "UTC" in text


def _seed_batch_run_for_download(tmp_path, monkeypatch) -> str:
    """POST a two-doc folder batch and return its run id. Cache dir is
    isolated so the operator's real gitignored config cannot make these
    brittle -- the pipeline falls back to the bundled sample."""
    import docx as docx_lib

    monkeypatch.setattr(app_module, "UPLOADS_DIR", tmp_path)
    cache_dir = tmp_path / "cache"
    monkeypatch.setattr(app_module, "AEM_CONFIG_DIR", cache_dir)
    monkeypatch.setattr(app_module, "SITE_CONFIG_PATH", cache_dir / "site.json")

    def make(title, slug):
        d = docx_lib.Document()
        d.add_paragraph().add_run(title).bold = True
        d.add_paragraph(f"Meta Title: {title}")
        d.add_paragraph("Meta Description: desc")
        d.add_paragraph("New Required URL: "
                        f"https://www.iffcotokio.co.in/health-insurance/{slug}")
        d.add_paragraph("Body.")
        buf = io.BytesIO()
        d.save(buf)
        buf.seek(0)
        return buf

    ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    files = [
        ("all_files", ("Articles/A1/one.docx", make("First Article", "first"), ct)),
        ("all_files", ("Articles/A2/two.docx", make("Second Article", "second"), ct)),
    ]
    resp = TestClient(app_module.app).post(
        "/run-batch", files=files, follow_redirects=False)
    return resp.headers["location"].rsplit("/", 1)[-1]



# ---------- package names carry the run's timestamp ----------

def test_package_stamp_is_compact_utc_derived_from_the_run():
    """Derived from the run's own created_at, not download time, so
    re-downloading a run still yields a byte-identical package."""
    from ui.app import _package_stamp
    assert _package_stamp("2026-09-16T04:07:33.583693+00:00") == "20260916-040733"
    # a record written by an older build, or a malformed value, must not
    # take the download route down -- the stamp is a convenience, not data
    assert _package_stamp("") == ""
    assert _package_stamp("not-a-date") == ""
    assert _package_stamp(None) == ""


def test_batch_package_filename_ends_with_the_run_timestamp(tmp_path, monkeypatch):
    run_id = _seed_batch_run_for_download(tmp_path, monkeypatch)
    response = TestClient(app_module.app).get(f"/batch-runs/{run_id}/package.zip")
    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    name = re.search(r'filename="([^"]+)\.zip"', disposition).group(1)
    assert re.search(r"-\d{8}-\d{6}$", name), name


def test_batch_package_manifest_name_matches_the_download_filename(tmp_path, monkeypatch):
    """AEM's Package Manager lists properties.xml's name, so the stamp has to
    reach it too -- otherwise every install overwrites the previous entry."""
    run_id = _seed_batch_run_for_download(tmp_path, monkeypatch)
    response = TestClient(app_module.app).get(f"/batch-runs/{run_id}/package.zip")
    name = re.search(r'filename="([^"]+)\.zip"',
                     response.headers["content-disposition"]).group(1)
    with zipfile.ZipFile(io.BytesIO(response.content)) as zf:
        props = zf.read("META-INF/vault/properties.xml").decode()
    assert f'<entry key="name">{name}</entry>' in props
