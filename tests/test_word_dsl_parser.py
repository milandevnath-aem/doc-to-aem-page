"""Regression checks for parser/word/dsl_parser.py and validators/word_structure.py.

Builds synthetic .docx files on the fly with python-docx (no Node/docx-js
dependency) so these run anywhere `pytest` runs.
"""

import base64
from pathlib import Path

import docx

from parser.word.dsl_parser import parse_docx
from validators.word_structure import validate_docx

# 1x1 transparent PNG -- smallest valid image python-docx can embed
PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _build_docx(path: Path, builder) -> str:
    document = docx.Document()
    builder(document)
    document.save(str(path))
    return str(path)


def _add_metadata(doc):
    doc.add_paragraph("[metadata]")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", "T"
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "D"


def _add_soft_broken(doc, *lines):
    """One paragraph holding several logical lines separated by soft line
    breaks (Shift+Enter) -- how the v2 client doc packs icon+title+body."""
    p = doc.add_paragraph()
    for i, line in enumerate(lines):
        run = p.add_run(line)
        if i < len(lines) - 1:
            run.add_break()
    return p


def _add_bullet(doc, text):
    p = doc.add_paragraph(text)
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = 1
    return p


def test_v1_v2_component_uses_only_htmltext_for_rte(tmp_path):
    """The RTE field on any component is exposed under one canonical name --
    `htmltext` -- and the source-authoring key (`description:` / `body:`)
    does NOT leak through to the output NCM."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: hero-banner]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Hero Title")
        doc.add_paragraph("description:")
        doc.add_paragraph("Hello world.")

    path = _build_docx(tmp_path / "hero.docx", build)
    hero = parse_docx(path)["blocks"][0]
    assert hero["fields"]["htmltext"] == "<p>Hello world.</p>"
    assert "description" not in hero["fields"]
    assert "body" not in hero["fields"]
    assert "headings" not in hero  # dropped -- was never populated for v2/v3 anyway


def test_v1_cardandrte_body_wins_when_both_body_and_description_authored(tmp_path):
    """A component authored with BOTH `body:` and `description:` keys -- body
    wins for the canonical htmltext (matches how AEM RTE components are
    typically named)."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: cardandRTE]")
        doc.add_paragraph("description:")
        doc.add_paragraph("Description text.")
        doc.add_paragraph("body:")
        doc.add_paragraph("Body text.")

    path = _build_docx(tmp_path / "both.docx", build)
    card = parse_docx(path)["blocks"][0]
    assert card["fields"]["htmltext"] == "<p>Body text.</p>"
    # neither source key survives the rename
    assert "body" not in card["fields"]
    assert "description" not in card["fields"]


def test_v2_lenient_item_uses_only_htmltext_for_rte(tmp_path):
    """v2 lenient items with an implicit description RTE field expose it as
    `htmltext` -- no separate description key on the output."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: featurecards]")
        doc.add_paragraph().add_run("Cards").bold = True
        _add_soft_broken(doc, "icon: location1.png", "Card Title", "Card body text.")

    path = _build_docx(tmp_path / "v2card.docx", build)
    item = parse_docx(path)["blocks"][0]["items"][0]
    assert item["fields"]["htmltext"] == "<p>Card body text.</p>"
    assert "description" not in item["fields"]


def test_component_without_rte_field_gets_no_htmltext(tmp_path):
    """A pure-scalar component (no description/body) must not gain an empty
    htmltext key."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: hero-banner]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Only a title.")

    path = _build_docx(tmp_path / "scalar.docx", build)
    hero = parse_docx(path)["blocks"][0]
    assert "htmltext" not in hero["fields"]


def test_parse_docx_extracts_rte_html_and_plain_fields(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "Test Page"
        table.cell(2, 0).text, table.cell(2, 1).text = "description", "A test page"

        doc.add_paragraph("[component: hero-banner]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Hero Title")
        doc.add_paragraph("description:")
        p = doc.add_paragraph()
        p.add_run("Plain text with ")
        p.add_run("bold text").bold = True
        p.add_run(" after it.")

    path = _build_docx(tmp_path / "sample.docx", build)
    result = parse_docx(path)

    assert result["page"]["properties"]["title"] == "Test Page"
    hero = result["blocks"][0]
    assert hero["name"] == "hero-banner"
    assert hero["fields"]["title"] == "Hero Title"
    assert hero["fields"]["htmltext"] == "<p>Plain text with <strong>bold text</strong> after it.</p>"


def test_parse_docx_table_under_rte_field_becomes_html_table(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        meta = doc.add_table(rows=3, cols=2)
        meta.cell(0, 0).text, meta.cell(0, 1).text = "Property", "Value"
        meta.cell(1, 0).text, meta.cell(1, 1).text = "title", "Test Page"
        meta.cell(2, 0).text, meta.cell(2, 1).text = "description", "A test page"

        doc.add_paragraph("[component: cardandRTE]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Motor Insurance by IFFCO Tokio - A Quick Look")
        doc.add_paragraph("description:")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Key Features", "Benefits"
        table.cell(1, 0).text, table.cell(1, 1).text = "Own Damage Cover", "Available"

    path = _build_docx(tmp_path / "sample.docx", build)
    result = parse_docx(path)

    card = result["blocks"][0]
    assert card["name"] == "cardandRTE"
    assert card["fields"]["htmltext"] == (
        "<table><tr><th>Key Features</th><th>Benefits</th></tr>"
        "<tr><td>Own Damage Cover</td><td>Available</td></tr></table>"
    )
    assert card["table"] is None  # table belonged to the RTE field, not the component


def test_parse_docx_lenient_icon_lines_split_into_items(tmp_path):
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: featurecards]")
        doc.add_paragraph().add_run("Benefits of Buying").bold = True
        _add_soft_broken(doc, "icon: location1.png", "Financial Protection",
                         "Insurance helps manage unexpected expenses.")
        _add_soft_broken(doc, "icon: location2.png", "Peace of Mind",
                         "You can focus on recovery.")

    path = _build_docx(tmp_path / "lenient.docx", build)
    result = parse_docx(path)

    comp = result["blocks"][0]
    assert comp["fields"]["title"] == "Benefits of Buying"
    assert len(comp["items"]) == 2
    first, second = comp["items"]
    assert first["fields"]["image"] == {"source": "reference", "name": "location1.png", "path": None}
    assert first["fields"]["title"] == "Financial Protection"
    assert first["fields"]["htmltext"] == "<p>Insurance helps manage unexpected expenses.</p>"
    assert second["fields"]["image"]["name"] == "location2.png"
    assert second["fields"]["title"] == "Peace of Mind"


def test_parse_docx_lenient_bold_labels_split_numbered_items(tmp_path):
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: cards-with-full-width||numbered]")
        doc.add_paragraph().add_run("Step-by-Step Guide").bold = True
        doc.add_paragraph().add_run("Step 1: Identify Your Insurance Needs").bold = True
        doc.add_paragraph("Determine why you need insurance.")
        _add_bullet(doc, "Health emergencies")
        _add_bullet(doc, "Vehicle accidents")
        doc.add_paragraph().add_run("Step 2: Determine the Coverage Amount").bold = True
        doc.add_paragraph("Choose adequate coverage.")

    path = _build_docx(tmp_path / "numbered.docx", build)
    result = parse_docx(path)

    comp = result["blocks"][0]
    assert comp["fields"]["title"] == "Step-by-Step Guide"
    assert len(comp["items"]) == 2
    step1, step2 = comp["items"]
    assert step1["fields"]["title"] == "Step 1: Identify Your Insurance Needs"
    assert step1["fields"]["htmltext"] == (
        "<p>Determine why you need insurance.</p>"
        "<ul><li>Health emergencies</li><li>Vehicle accidents</li></ul>"
    )
    assert step2["fields"]["title"] == "Step 2: Determine the Coverage Amount"
    assert step2["fields"]["htmltext"] == "<p>Choose adequate coverage.</p>"


def test_parse_docx_lenient_hero_implicit_title_and_inline_keys(tmp_path):
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: hero-banner]")
        doc.add_paragraph("Tips for Buying Insurance: A Complete Guide")
        doc.add_paragraph("An intro paragraph under the hero.")
        _add_soft_broken(doc, "Data-of-release: June 23, 2026", "Estimated-read-time: 4 min read")

    path = _build_docx(tmp_path / "hero.docx", build)
    result = parse_docx(path)

    hero = result["blocks"][0]
    assert hero["fields"]["title"] == "Tips for Buying Insurance: A Complete Guide"
    assert hero["fields"]["htmltext"] == "<p>An intro paragraph under the hero.</p>"
    assert hero["fields"]["data-of-release"] == "June 23, 2026"
    assert hero["fields"]["estimated-read-time"] == "4 min read"


def test_parse_docx_strict_keys_still_win_over_lenient_rules(tmp_path):
    # strict-format doc: explicit keys must behave exactly as before, no
    # implicit items or titles
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: hero-banner]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Explicit Title")
        doc.add_paragraph("description:")
        p = doc.add_paragraph()
        p.add_run("All bold description line.").bold = True

    path = _build_docx(tmp_path / "strict.docx", build)
    result = parse_docx(path)

    hero = result["blocks"][0]
    assert hero["fields"]["title"] == "Explicit Title"
    # bold paragraph under an active RTE key is content, not an item delimiter
    assert hero["fields"]["htmltext"] == "<p><strong>All bold description line.</strong></p>"
    assert hero["items"] == []


def test_parse_docx_extracts_embedded_image_to_assets_dir(tmp_path):
    img_path = tmp_path / "pic.png"
    img_path.write_bytes(PNG_1PX)

    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: featurecards]")
        doc.add_paragraph().add_run("Cards With Pasted Images").bold = True
        doc.add_picture(str(img_path))
        doc.add_paragraph("Pasted Card Title")
        doc.add_paragraph("Pasted card description.")

    path = _build_docx(tmp_path / "embedded.docx", build)
    assets_dir = tmp_path / "assets"
    result = parse_docx(path, assets_dir=str(assets_dir))

    comp = result["blocks"][0]
    assert len(comp["items"]) == 1
    item = comp["items"][0]
    image = item["fields"]["image"]
    assert image["source"] == "embedded"
    assert image["name"] == "image-1.png"
    assert image["path"] == (assets_dir / "image-1.png").as_posix()
    assert (assets_dir / "image-1.png").read_bytes() == PNG_1PX
    assert item["fields"]["title"] == "Pasted Card Title"
    assert item["fields"]["htmltext"] == "<p>Pasted card description.</p>"


def test_parse_docx_embedded_image_without_assets_dir_records_no_path(tmp_path):
    img_path = tmp_path / "pic.png"
    img_path.write_bytes(PNG_1PX)

    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: featurecards]")
        doc.add_paragraph().add_run("Cards").bold = True
        doc.add_picture(str(img_path))
        doc.add_paragraph("Card Title")

    path = _build_docx(tmp_path / "no_assets.docx", build)
    result = parse_docx(path)

    image = result["blocks"][0]["items"][0]["fields"]["image"]
    assert image == {"source": "embedded", "name": "image-1.png", "path": None}


def test_validate_docx_flags_unrecognized_marker_and_missing_metadata(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "Only Title"
        doc.add_paragraph("[component hero-banner]")  # missing colon -> unrecognized marker
        doc.add_paragraph("title:")
        doc.add_paragraph("dropped content")

    path = _build_docx(tmp_path / "broken.docx", build)
    issues = validate_docx(path)
    codes = {i.code for i in issues}

    assert "unrecognized-marker" in codes
    assert "missing-required-metadata" in codes  # description missing
    assert all(i.severity == "error" for i in issues)


def test_validate_docx_flags_card_outside_component(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "T"
        table.cell(2, 0).text, table.cell(2, 1).text = "description", "D"
        doc.add_paragraph("[card: 01]")  # no active component

    path = _build_docx(tmp_path / "orphan_card.docx", build)
    issues = validate_docx(path)

    assert any(i.code == "card-outside-component" for i in issues)


def test_validate_docx_accepts_teasercomponent_and_lenient_content(tmp_path):
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: teasercomponent]")
        doc.add_paragraph("Expert Tips for Buying Insurance")
        _add_soft_broken(doc, "icon: location6.png", "Expert Tips", "Always assess your needs.")

    path = _build_docx(tmp_path / "teaser.docx", build)
    issues = validate_docx(path)

    assert not any(i.severity == "error" for i in issues)
    # lenient content counts as content -- no empty-component warning either
    assert not any(i.code == "empty-component" for i in issues)


def test_parse_real_v2_client_doc_end_to_end():
    """Regression against the actual client-authored v2 sample: lenient
    authoring (component markers, no field keys) must capture all content."""
    fixture = Path(__file__).parent / "fixtures" / "article-v2-sample.docx"

    issues = validate_docx(str(fixture))
    assert issues == []

    result = parse_docx(str(fixture))
    blocks = result["blocks"]
    assert [(b["name"], b["variant"]) for b in blocks] == [
        ("hero-banner", None),
        ("cards-with-full-width", "icon"),
        ("cards-with-full-width", "numbered"),
        ("featurecards", None),
        ("featurecards", None),
        ("featurecards", None),
        ("teasercomponent", None),
    ]

    hero = blocks[0]
    assert hero["fields"]["title"].startswith("Tips for Buying Insurance")
    assert hero["fields"]["data-of-release"] == "June 23, 2026"

    assert len(blocks[1]["items"]) == 3   # icon cards
    steps = blocks[2]["items"]
    assert len(steps) == 7                # Step 1..7
    assert steps[0]["fields"]["title"] == "Step 1: Identify Your Insurance Needs"
    assert [len(b["items"]) for b in blocks[3:6]] == [5, 7, 6]  # feature card groups
    assert blocks[3]["items"][0]["fields"]["image"]["name"] == "location1.png"

    teaser = blocks[6]["items"][0]
    assert teaser["fields"]["title"] == "Expert Tips for Buying Insurance"
    assert teaser["fields"]["htmltext"].count("<li>") == 6

    assert result["page"]["properties"]["title"] == "Family Health Insurance | IFFCO Tokio"


def test_validate_docx_clean_doc_has_no_errors(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
        table.cell(1, 0).text, table.cell(1, 1).text = "title", "T"
        table.cell(2, 0).text, table.cell(2, 1).text = "description", "D"
        doc.add_paragraph("[component: hero-banner]")
        doc.add_paragraph("title:")
        doc.add_paragraph("Hero Title")

    path = _build_docx(tmp_path / "clean.docx", build)
    issues = validate_docx(path)

    assert not any(i.severity == "error" for i in issues)


def test_new_hero_banner_extracts_banner_image_filename(tmp_path):
    """The `Banner Image: <file>` line inside a `[component:new-hero-banner]`
    block is a component-specific asset marker: filename lands on the
    component's `imageFile` field, and does NOT leak into `htmltext` or
    the implicit `title` (which the mapper later overrides with meta title
    anyway)."""
    def build(doc):
        _add_metadata(doc)
        _add_soft_broken(
            doc,
            "[component:new-hero-banner]",
            "Banner Image: car-insurance.jpg",
            "",
            "A family health insurance plan helps pay hospital bills.",
        )

    path = _build_docx(tmp_path / "hero_asset.docx", build)
    hero = parse_docx(path)["blocks"][0]

    assert hero["name"] == "new-hero-banner"
    assert hero["fields"]["imageFile"] == "car-insurance.jpg"
    # description survives, without the marker line polluting it
    assert "Banner Image" not in hero["fields"].get("htmltext", "")
    assert "family health insurance" in hero["fields"].get("htmltext", "")
    # implicit title must NOT be the banner-image line
    assert hero["fields"].get("title", "") != "Banner Image: car-insurance.jpg"


def test_new_hero_banner_banner_image_marker_case_insensitive(tmp_path):
    """Filename lookup is case-insensitive; marker keyword tolerates casing
    variance (`banner image`, `BANNER IMAGE`)."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("BANNER IMAGE: HERO-BANNER.PNG")
        doc.add_paragraph("Body copy.")

    path = _build_docx(tmp_path / "hero_case.docx", build)
    hero = parse_docx(path)["blocks"][0]

    assert hero["fields"]["imageFile"] == "HERO-BANNER.PNG"
    assert "BANNER IMAGE" not in hero["fields"].get("htmltext", "")


def test_banner_image_marker_ignored_outside_new_hero_banner(tmp_path):
    """The `Banner Image:` shortcut only applies inside a
    `[component:new-hero-banner]` block. Any other component treats the
    same line as ordinary prose so we don't hijack unrelated authoring."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:page-description]")
        doc.add_paragraph("Banner Image: not-a-banner.jpg")
        doc.add_paragraph("Content.")

    path = _build_docx(tmp_path / "other.docx", build)
    block = parse_docx(path)["blocks"][0]

    assert "imageFile" not in block["fields"]


def test_faq_heading_splits_into_own_faq_block(tmp_path):
    """Real client docs end the article body with a bold FAQ heading
    (`FAQs`, `FAQS`, `Frequently Asked Questions (FAQ)`) followed by
    Q/A pairs. That FAQ section must land in its own structured `faq`
    block, not fold into the main body page-description -- otherwise
    Q/A copy pollutes the article body content on the rendered page."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner Image: home.jpg")
        doc.add_paragraph("Banner intro.")
        # section title as fully-bold [normal] paragraph
        p = doc.add_paragraph()
        p.add_run("Family Health Insurance").bold = True
        doc.add_paragraph("Main body copy.")
        # FAQ section
        p = doc.add_paragraph()
        p.add_run("FAQs").bold = True
        p = doc.add_paragraph()
        p.add_run("Q1?").bold = True
        doc.add_paragraph("A1.")
        p = doc.add_paragraph()
        p.add_run("Q2?").bold = True
        doc.add_paragraph("A2.")

    path = _build_docx(tmp_path / "faq-split.docx", build)
    blocks = parse_docx(path)["blocks"]

    # Three blocks: hero + main body page-description + structured faq
    names = [b["name"] for b in blocks]
    assert names == [
        "new-hero-banner",
        "page-description",   # main body
        "faq",
    ], names

    main = blocks[1]["fields"].get("htmltext", "")
    faq = blocks[2]

    # main body has the section title + body copy, NO FAQ
    assert "Family Health Insurance" in main
    assert "Main body copy" in main
    assert "FAQs" not in main
    assert "Q1?" not in main
    assert "A1." not in main

    # FAQ block carries the FAQ heading as its title, and each question as
    # its own item with the answer as that item's htmltext
    assert faq["fields"]["title"] == "FAQs"
    items = faq["items"]
    assert [i["fields"]["question"] for i in items] == ["Q1?", "Q2?"]
    assert items[0]["fields"]["htmltext"] == "<p>A1.</p>"
    assert items[1]["fields"]["htmltext"] == "<p>A2.</p>"


def test_frequently_asked_questions_label_also_triggers_faq_split(tmp_path):
    """The FAQ trigger accepts the long form too, with optional
    parenthetical `(FAQ)` suffix and any casing."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner Image: home.jpg")
        doc.add_paragraph("Body intro.")
        p = doc.add_paragraph()
        p.add_run("Section Two").bold = True
        doc.add_paragraph("More body copy.")
        p = doc.add_paragraph()
        p.add_run("Frequently Asked Questions (FAQ)").bold = True
        p = doc.add_paragraph()
        p.add_run("Q?").bold = True
        doc.add_paragraph("A.")

    path = _build_docx(tmp_path / "faq-long.docx", build)
    blocks = parse_docx(path)["blocks"]
    assert [b["name"] for b in blocks][-1] == "faq"
    faq = blocks[-1]
    assert faq["fields"]["title"] == "Frequently Asked Questions (FAQ)"
    assert faq["items"][0]["fields"]["question"] == "Q?"
    assert faq["items"][0]["fields"]["htmltext"] == "<p>A.</p>"


def test_new_hero_banner_closed_on_empty_paragraph_gap(tmp_path):
    """Some real client docs place a horizontal rule between the banner
    section and the article body, with no bold or Heading paragraph
    that acts as a section title. The .docx XML shows all border
    values `nil` so the rule itself can't be detected -- but an EMPTY
    paragraph always sits at that boundary. Close the hero on the
    empty-paragraph gap once the hero has accumulated real content."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner Image: home.jpg")
        doc.add_paragraph("Banner intro paragraph one.")
        doc.add_paragraph("Banner intro paragraph two.")
        doc.add_paragraph()  # empty separator (visually a HR)
        doc.add_paragraph("Body copy paragraph starts here.")
        doc.add_paragraph("Another body paragraph.")

    path = _build_docx(tmp_path / "empty-gap.docx", build)
    blocks = parse_docx(path)["blocks"]

    assert [b["name"] for b in blocks] == ["new-hero-banner", "page-description"]

    hero_html = blocks[0]["fields"].get("htmltext", "")
    body_html = blocks[1]["fields"].get("htmltext", "")

    assert "Banner intro paragraph one" in hero_html
    assert "Banner intro paragraph two" in hero_html
    assert "Body copy" not in hero_html

    assert "Body copy paragraph starts here" in body_html
    assert "Another body paragraph" in body_html


def test_new_hero_banner_closed_on_first_bold_paragraph_after_intro(tmp_path):
    """Not every real client doc uses a Word Heading style for its
    section titles -- several use a fully-bold `[normal]` paragraph
    instead. Close the hero on that too, once at least one prose
    paragraph has already accumulated as the banner description."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner Image: home.jpg")
        doc.add_paragraph("Banner intro line one.")
        # section title as bold-only [normal] paragraph
        p = doc.add_paragraph()
        p.add_run("Family Health Insurance").bold = True
        doc.add_paragraph("Body copy under the bold section title.")
        doc.add_paragraph("Another body paragraph.")

    path = _build_docx(tmp_path / "bold-close.docx", build)
    blocks = parse_docx(path)["blocks"]

    assert [b["name"] for b in blocks] == ["new-hero-banner", "page-description"]

    hero_html = blocks[0]["fields"].get("htmltext", "")
    body_html = blocks[1]["fields"].get("htmltext", "")

    # banner description carries the intro paragraph only
    assert "Banner intro line one" in hero_html
    assert "Family Health Insurance" not in hero_html
    assert "Body copy" not in hero_html

    # implicit page-description carries the bold section title +
    # everything after it
    assert "Family Health Insurance" in body_html
    assert "Body copy under the bold section title" in body_html


def test_new_hero_banner_closed_on_first_heading_after_its_prose(tmp_path):
    """Real client docs use a SINGLE `[component:new-hero-banner]` marker
    then leave the article body unmarked (no `[component:page-description]`
    to switch to). A Word Heading paragraph is the natural end of the
    banner section -- close the hero and start an implicit
    page-description so the rest of the doc lands in a separate
    component instead of being folded into bannerDescription."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner Image: car.jpg")
        doc.add_paragraph("Banner intro prose one.")
        doc.add_paragraph("Banner intro prose two.")
        doc.add_paragraph()  # empty separator
        doc.add_paragraph("What is Cashless Treatment?", style="Heading 2")
        doc.add_paragraph("Body copy under the H2.")
        doc.add_paragraph("Another body paragraph.")

    path = _build_docx(tmp_path / "auto-close.docx", build)
    blocks = parse_docx(path)["blocks"]

    # Two blocks: hero-banner + implicit page-description
    assert [b["name"] for b in blocks] == ["new-hero-banner", "page-description"]

    hero = blocks[0]
    assert hero["fields"]["imageFile"] == "car.jpg"
    hero_html = hero["fields"].get("htmltext", "")
    # banner description holds only the intro paragraphs -- headings
    # and body copy that follow are NOT in it
    assert "Banner intro prose one" in hero_html
    assert "Banner intro prose two" in hero_html
    assert "Cashless Treatment" not in hero_html
    assert "Body copy" not in hero_html

    body = blocks[1]
    body_html = body["fields"].get("htmltext", "")
    # implicit page-description carries the H2 + body copy
    assert "Cashless Treatment" in body_html
    assert "Body copy under the H2." in body_html
    assert "Another body paragraph" in body_html


def test_url_captured_from_any_label_containing_url(tmp_path):
    """Real client docs spell the URL key four different ways. Capture is
    driven by the value's shape, not a fixed list of labels."""
    def build(doc):
        doc.add_paragraph("URL Need to be Careated:https://www.iffcotokio.co.in/health-insurance/why-family")
        doc.add_paragraph("Meta Title: Why Family Health Insurance")
        doc.add_paragraph("Meta Description: D")
        doc.add_paragraph("Keywords:   family health insurance, tax")
        doc.add_paragraph("[component: new-hero-banner]")
        doc.add_paragraph("Banner Image: hero.jpg")
        doc.add_paragraph("Some banner prose.")

    path = _build_docx(tmp_path / "a4.docx", build)
    props = parse_docx(path, site_origin="https://www.iffcotokio.co.in")["page"]["properties"]
    assert props["url"] == "https://www.iffcotokio.co.in/health-insurance/why-family"
    assert props["keywords"] == "family health insurance, tax"
    assert props["title"] == "Why Family Health Insurance"


def test_old_url_label_routes_to_old_url_not_url(tmp_path):
    def build(doc):
        doc.add_paragraph("Old URL: https://www.iffcotokio.co.in/old/page")
        doc.add_paragraph("New Required URL : https://www.iffcotokio.co.in/new/page")
        doc.add_paragraph("Meta Title: T")
        doc.add_paragraph("Meta Description: D")
        doc.add_paragraph("[component: new-hero-banner]")
        doc.add_paragraph("Prose.")

    path = _build_docx(tmp_path / "urls.docx", build)
    props = parse_docx(path, site_origin="https://www.iffcotokio.co.in")["page"]["properties"]
    assert props["old_url"] == "https://www.iffcotokio.co.in/old/page"
    assert props["url"] == "https://www.iffcotokio.co.in/new/page"


def test_url_label_with_off_origin_value_is_not_captured(tmp_path):
    """A URL on another domain is prose, not this page's URL."""
    def build(doc):
        doc.add_paragraph("Reference URL: https://example.com/somewhere")
        doc.add_paragraph("Meta Title: T")
        doc.add_paragraph("Meta Description: D")
        doc.add_paragraph("[component: new-hero-banner]")
        doc.add_paragraph("Prose.")

    path = _build_docx(tmp_path / "off.docx", build)
    props = parse_docx(path, site_origin="https://www.iffcotokio.co.in")["page"]["properties"]
    assert "url" not in props


def test_faq_section_becomes_structured_faq_block(tmp_path):
    """A bold FAQ label opens a faq block; each bold line is a question and
    the plain paragraphs under it are that question's answer."""
    def build(doc):
        doc.add_paragraph("Meta Title: T")
        doc.add_paragraph("Meta Description: D")
        doc.add_paragraph("[component: new-hero-banner]")
        doc.add_paragraph("Banner prose.")
        doc.add_paragraph().add_run("Body Heading").bold = True
        doc.add_paragraph("Article body paragraph.")
        doc.add_paragraph().add_run("FAQs").bold = True
        doc.add_paragraph().add_run("Can parents be included?").bold = True
        doc.add_paragraph("Yes, parents can be included.")
        doc.add_paragraph("Subject to age limits.")
        doc.add_paragraph().add_run("Is it mandatory?").bold = True
        doc.add_paragraph("No, it is not mandatory.")

    path = _build_docx(tmp_path / "faq.docx", build)
    blocks = parse_docx(path)["blocks"]
    faqs = [b for b in blocks if b["name"] == "faq"]
    assert len(faqs) == 1
    items = faqs[0]["items"]
    assert [i["fields"]["question"] for i in items] == [
        "Can parents be included?", "Is it mandatory?"]
    assert items[0]["fields"]["htmltext"] == (
        "<p>Yes, parents can be included.</p><p>Subject to age limits.</p>")
    assert items[1]["fields"]["htmltext"] == "<p>No, it is not mandatory.</p>"


def test_faq_answers_keep_bullets_as_list_markup(tmp_path):
    def build(doc):
        doc.add_paragraph("Meta Title: T")
        doc.add_paragraph("Meta Description: D")
        doc.add_paragraph("[component: new-hero-banner]")
        doc.add_paragraph("Banner prose.")
        doc.add_paragraph().add_run("FAQs").bold = True
        doc.add_paragraph().add_run("What is covered?").bold = True
        _add_bullet(doc, "Hospitalisation")
        _add_bullet(doc, "Day care")

    path = _build_docx(tmp_path / "faqbul.docx", build)
    faq = [b for b in parse_docx(path)["blocks"] if b["name"] == "faq"][0]
    assert faq["items"][0]["fields"]["htmltext"] == (
        "<ul><li>Hospitalisation</li><li>Day care</li></ul>")


def test_faq_label_with_no_questions_yields_empty_faq_block(tmp_path):
    """A FAQ heading with nothing under it must not silently vanish, and must
    not swallow the block into the article body either."""
    def build(doc):
        doc.add_paragraph("Meta Title: T")
        doc.add_paragraph("Meta Description: D")
        doc.add_paragraph("[component: new-hero-banner]")
        doc.add_paragraph("Banner prose.")
        doc.add_paragraph().add_run("Frequently Asked Questions").bold = True

    path = _build_docx(tmp_path / "faqempty.docx", build)
    faqs = [b for b in parse_docx(path)["blocks"] if b["name"] == "faq"]
    assert len(faqs) == 1
    assert faqs[0]["items"] == []


def test_validate_docx_warns_when_faq_label_has_no_questions(tmp_path):
    """Spec's error-handling table: an FAQ label with no bold question under
    it must warn (the page still builds, faq_count=0). `_start_faq` fills the
    block's `title`, so `empty-component` is satisfied and never fires -- this
    is the only signal the operator gets."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: page-description]")
        doc.add_paragraph("description:")
        doc.add_paragraph("Body prose.")
        doc.add_paragraph().add_run("FAQs").bold = True
        doc.add_paragraph("Plain prose where a bold question should be.")

    issues = validate_docx(_build_docx(tmp_path / "faq-no-questions.docx", build))

    faq_issues = [i for i in issues if i.code == "faq-without-questions"]
    assert len(faq_issues) == 1
    assert faq_issues[0].severity == "warning"
    assert "FAQs" in faq_issues[0].message
    assert not any(i.severity == "error" for i in issues)   # page still builds


def test_validate_docx_does_not_warn_when_faq_has_questions(tmp_path):
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component: page-description]")
        doc.add_paragraph("description:")
        doc.add_paragraph("Body prose.")
        doc.add_paragraph().add_run("FAQs").bold = True
        doc.add_paragraph().add_run("Is it covered?").bold = True
        doc.add_paragraph("Yes.")

    issues = validate_docx(_build_docx(tmp_path / "faq-ok.docx", build))

    assert not any(i.code == "faq-without-questions" for i in issues)


# ---- city-page authoring (v2 client docs) ----

ORIGIN = "https://www.iffcotokio.co.in"


def test_meta_value_on_the_line_below_its_label_is_captured(tmp_path):
    """City docs put a meta label and its value on either side of a soft
    line break:

        Meta Title:
        Health Insurance in New Delhi | IFFCO-TOKIO

    The label line matched with an empty value, so nothing was stored and
    the value line -- arriving before any component -- was dropped. The
    page then shipped with no jcr:title and no jcr:description."""
    def build(doc):
        doc.add_paragraph("Health Insurance in New Delhi", style="Heading 1")
        doc.add_paragraph(f"New Required URL: {ORIGIN}/health-insurance/city/new-delhi")
        _add_soft_broken(doc, "Meta Title:", "Health Insurance in New Delhi | IFFCO-TOKIO")
        _add_soft_broken(doc, "Meta Description:", "Know what to check about coverage.")
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner prose.")

    path = _build_docx(tmp_path / "city-meta.docx", build)
    props = parse_docx(path, site_origin=ORIGIN)["page"]["properties"]

    assert props["title"] == "Health Insurance in New Delhi | IFFCO-TOKIO"
    assert props["description"] == "Know what to check about coverage."
    assert props["url"] == f"{ORIGIN}/health-insurance/city/new-delhi"


def test_bare_banner_image_label_does_not_become_the_hero_title(tmp_path):
    """v2 city docs carry `Banner Image:` with no filename -- the image
    comes from elsewhere. The bare label must be consumed, not promoted to
    the hero's implicit title."""
    def build(doc):
        _add_metadata(doc)
        _add_soft_broken(
            doc,
            "[component:new-hero-banner]",
            "Banner Image:",
            "",
            "A medical emergency can bring an unexpected expense.",
        )

    path = _build_docx(tmp_path / "bare-banner.docx", build)
    hero = parse_docx(path)["blocks"][0]

    assert hero["name"] == "new-hero-banner"
    assert hero["fields"].get("title", "") != "Banner Image:"
    assert "imageFile" not in hero["fields"]
    assert "medical emergency" in hero["fields"].get("htmltext", "")
    assert "Banner Image" not in hero["fields"].get("htmltext", "")


def test_heading_styled_faq_label_and_questions_build_a_faq_block(tmp_path):
    """v2 city docs style the FAQ label as Heading 1 and each question as
    Heading 3 -- neither is bold. FAQ detection required bold, so the whole
    Q&A section folded into page-description and no Content Fragments were
    generated for it."""
    def build(doc):
        _add_metadata(doc)
        doc.add_paragraph("[component:new-hero-banner]")
        doc.add_paragraph("Banner prose.")
        doc.add_paragraph("Conclusion", style="Heading 1")
        doc.add_paragraph("Closing body copy.")
        doc.add_paragraph("FAQs", style="Heading 1")
        doc.add_paragraph("Is health insurance useful here?", style="Heading 3")
        doc.add_paragraph("It can help manage eligible medical expenses.")
        doc.add_paragraph("Can I buy online?", style="Heading 3")
        doc.add_paragraph("Yes, depending on the product.")

    path = _build_docx(tmp_path / "city-faq.docx", build)
    blocks = parse_docx(path)["blocks"]
    names = [b["name"] for b in blocks]

    assert "faq" in names, names
    faq = next(b for b in blocks if b["name"] == "faq")
    assert [i["fields"]["question"] for i in faq["items"]] == [
        "Is health insurance useful here?",
        "Can I buy online?",
    ]
    assert "eligible medical expenses" in faq["items"][0]["fields"].get("htmltext", "")

    # and the Q&A must NOT also sit in the article body
    body = " ".join(b["fields"].get("htmltext", "") for b in blocks
                    if b["name"] == "page-description")
    assert "Is health insurance useful here?" not in body
