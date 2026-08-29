"""V3 markerless article parser."""
from pathlib import Path

import docx

from parser.word.dsl_parser import parse_docx


def _bold(p, text):
    r = p.add_run(text)
    r.bold = True
    return r


def _bold_para(doc, text):
    p = doc.add_paragraph()
    _bold(p, text)
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(text)
    numPr = p._p.get_or_add_pPr().get_or_add_numPr()
    numPr.get_or_add_ilvl().val = 0
    numPr.get_or_add_numId().val = 1
    return p


def _heading(doc, text, level=2):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def _build(path, builder):
    document = docx.Document()
    builder(document)
    document.save(str(path))
    return str(path)


def _write_v3_metadata_header(doc, h1="Sample H1", include_old_url=True):
    _bold_para(doc, h1)
    if include_old_url:
        doc.add_paragraph("Old URL: https://old.example.com/x")
    doc.add_paragraph("New Required URL: https://new.example.com/x")
    doc.add_paragraph("Meta Title: The Title")
    doc.add_paragraph("Meta Description: The description.")
    doc.add_paragraph("Meta keywords: kw1, kw2, kw3")


def test_v3_metadata_block_captures_all_five_keys(tmp_path):
    def build(doc):
        _write_v3_metadata_header(doc)
        doc.add_paragraph("Body content starts here.")

    path = _build(tmp_path / "meta.docx", build)
    result = parse_docx(path)

    props = result["page"]["properties"]
    assert result["format"] == "v3"
    assert props["h1"] == "Sample H1"
    assert props["old_url"] == "https://old.example.com/x"
    assert props["url"] == "https://new.example.com/x"
    assert props["title"] == "The Title"
    assert props["description"] == "The description."
    assert props["keywords"] == "kw1, kw2, kw3"


def test_v3_missing_meta_keys_become_empty_string(tmp_path):
    def build(doc):
        _write_v3_metadata_header(doc, include_old_url=False)
        doc.add_paragraph("Body.")

    path = _build(tmp_path / "meta-partial.docx", build)
    props = parse_docx(path)["page"]["properties"]
    assert props["old_url"] == ""
    assert props["url"].startswith("https://")


def test_v3_body_becomes_single_page_description_with_headings_and_bullets(tmp_path):
    def build(doc):
        _write_v3_metadata_header(doc)
        doc.add_paragraph("Intro paragraph.")
        _heading(doc, "Section Heading", 2)
        doc.add_paragraph("Section text.")
        _bullet(doc, "point one")
        _bullet(doc, "point two")

    path = _build(tmp_path / "body.docx", build)
    result = parse_docx(path)
    blocks = result["blocks"]
    assert len(blocks) == 1
    body = blocks[0]
    assert body["name"] == "page-description"
    assert body["fields"]["body"] == (
        "<p>Intro paragraph.</p>"
        "<h2>Section Heading</h2>"
        "<p>Section text.</p>"
        "<ul><li>point one</li><li>point two</li></ul>"
    )


def test_v3_table_in_body_becomes_rte_table(tmp_path):
    def build(doc):
        _write_v3_metadata_header(doc)
        doc.add_paragraph("Compare:")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text, tbl.cell(0, 1).text = "A", "B"
        tbl.cell(1, 0).text, tbl.cell(1, 1).text = "1", "2"

    path = _build(tmp_path / "table.docx", build)
    body_html = parse_docx(path)["blocks"][0]["fields"]["body"]
    assert "<p>Compare:</p>" in body_html
    assert "<table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table>" in body_html


def test_v3_faq_section_splits_into_question_answer_items(tmp_path):
    def build(doc):
        _write_v3_metadata_header(doc)
        doc.add_paragraph("Article body.")
        _bold_para(doc, "FAQS")
        _bold_para(doc, "What is pre-authorization?")
        doc.add_paragraph("Pre-authorization is an approval...")
        _bold_para(doc, "Is it required for both claim types?")
        doc.add_paragraph("Only for cashless treatment.")
        _bullet(doc, "Note one")

    path = _build(tmp_path / "faq.docx", build)
    blocks = parse_docx(path)["blocks"]
    assert [b["name"] for b in blocks] == ["page-description", "faq"]
    faq = blocks[1]
    assert faq["fields"]["title"] == "FAQS"
    assert len(faq["items"]) == 2
    q1, q2 = faq["items"]
    assert q1["fields"]["question"] == "What is pre-authorization?"
    assert q1["fields"]["answer"] == "<p>Pre-authorization is an approval...</p>"
    assert q2["fields"]["question"] == "Is it required for both claim types?"
    assert q2["fields"]["answer"] == "<p>Only for cashless treatment.</p><ul><li>Note one</li></ul>"


def test_v3_final_thought_splits_out_trailing_page_description(tmp_path):
    def build(doc):
        _write_v3_metadata_header(doc)
        doc.add_paragraph("Body.")
        _bold_para(doc, "FAQS")
        _bold_para(doc, "Q?")
        doc.add_paragraph("A.")
        _bold_para(doc, "Final Thought")
        doc.add_paragraph("Closing paragraph.")

    path = _build(tmp_path / "final.docx", build)
    blocks = parse_docx(path)["blocks"]
    assert [b["name"] for b in blocks] == ["page-description", "faq", "page-description"]
    closing = blocks[2]["fields"]["body"]
    assert closing == "<h2>Final Thought</h2><p>Closing paragraph.</p>"


def test_v3_real_client_doc_e2e():
    fixture = Path(__file__).parent / "fixtures" / "article-v3-sample.docx"
    result = parse_docx(str(fixture))

    assert result["format"] == "v3"
    props = result["page"]["properties"]
    assert props["h1"] == "Cashless Treatment vs Reimbursement Claims"
    assert props["old_url"] == ""  # missing in the doc
    assert props["url"].startswith("https://www.iffcotokio.co.in/")
    assert props["title"] == "Cashless vs Reimbursement Health Insurance Claims"
    assert props["description"].startswith("Learn about the cashless")
    assert "health insurance claim process" in props["keywords"]

    names = [b["name"] for b in result["blocks"]]
    assert names == ["page-description", "faq", "page-description"]

    faq_items = result["blocks"][1]["items"]
    assert len(faq_items) == 8
    assert faq_items[0]["fields"]["question"] == "What is pre-authorization in health insurance?"

    body_html = result["blocks"][0]["fields"]["body"]
    assert "<table>" in body_html                            # differences table
    assert "<h3>What is Cashless Treatment?</h3>" in body_html
    assert "<ul><li>" in body_html                            # bullet list

    closing = result["blocks"][2]["fields"]["body"]
    assert "<h2>Final Thought</h2>" in closing


def test_v1_v2_still_works_and_format_annotates_output(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        tbl = doc.add_table(rows=3, cols=2)
        tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Property", "Value"
        tbl.cell(1, 0).text, tbl.cell(1, 1).text = "title", "T"
        tbl.cell(2, 0).text, tbl.cell(2, 1).text = "description", "D"

    path = _build(tmp_path / "strict.docx", build)
    result = parse_docx(path)
    assert result["format"] == "v1v2"
    assert result["page"]["properties"]["title"] == "T"


def test_parse_docx_format_override_forces_path(tmp_path):
    """Explicit format=v3 must win over auto-detection."""
    def build(doc):
        _write_v3_metadata_header(doc)
        doc.add_paragraph("Body.")

    path = _build(tmp_path / "override.docx", build)
    assert parse_docx(path, format="v3")["format"] == "v3"
    # forcing v1v2 on a markerless doc: still parses (validator would complain)
    result = parse_docx(path, format="v1v2")
    assert result["format"] == "v1v2"
