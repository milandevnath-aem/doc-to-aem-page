"""Format detection between the marker-based (v1/v2) and markerless (v3)
authoring styles.
"""

from pathlib import Path

import docx

from parser.word.format_detect import detect_format


def _build(path, builder):
    document = docx.Document()
    builder(document)
    document.save(str(path))
    return docx.Document(str(path))


def test_detect_marker_based_returns_v1v2(tmp_path):
    def build(doc):
        doc.add_paragraph("[metadata]")
        doc.add_paragraph("[component: hero-banner]")

    document = _build(tmp_path / "marker.docx", build)
    assert detect_format(document) == "v1v2"


def test_detect_markerless_with_meta_keys_returns_v3(tmp_path):
    def build(doc):
        doc.add_paragraph().add_run("Some Page H1").bold = True
        doc.add_paragraph("Meta Title: Something")
        doc.add_paragraph("Meta Description: A description")

    document = _build(tmp_path / "v3.docx", build)
    assert detect_format(document) == "v3"


def test_detect_blank_doc_defaults_to_v1v2(tmp_path):
    document = _build(tmp_path / "empty.docx", lambda doc: None)
    assert detect_format(document) == "v1v2"


def test_detect_real_v3_fixture():
    fixture = Path(__file__).parent / "fixtures" / "article-v3-sample.docx"
    document = docx.Document(str(fixture))
    assert detect_format(document) == "v3"


def test_detect_real_v2_fixture():
    fixture = Path(__file__).parent / "fixtures" / "article-v2-sample.docx"
    document = docx.Document(str(fixture))
    assert detect_format(document) == "v1v2"
