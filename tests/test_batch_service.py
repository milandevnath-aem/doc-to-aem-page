"""Batch mode -- one folder, many docs, one pipeline run per docx.

Each subfolder = one page. The subfolder NAME does not leak into the JCR
path -- it is only a container for a page's docx + its assets. Page path
comes from pageNameStrategy applied to doc contents, same as single-doc.
"""

from pathlib import Path

import docx

from aem_client.config import FaqConfig, PageTypeConfig
from ui.batch_service import run_batch


def _sample_config() -> PageTypeConfig:
    return PageTypeConfig(
        type="article",
        template_path="/conf/x",
        content_root="/content/site/articles",
        dam_root="/content/dam/site/articles",
        page_name_strategy="slug_from_meta_title",
        component_mappings={"page-description": "site/components/text"},
    )


def _write_v3_doc(path: Path, title: str, body: str = "Body content.") -> None:
    document = docx.Document()
    p = document.add_paragraph()
    p.add_run(title).bold = True
    document.add_paragraph(f"Meta Title: {title}")
    document.add_paragraph("Meta Description: Some description.")
    document.add_paragraph("Meta keywords: kw")
    document.add_paragraph(body)
    document.save(str(path))


def test_batch_runs_pipeline_per_file_and_returns_ordered_entries(tmp_path):
    for i, (folder, title) in enumerate([
        ("Article1", "First Article"),
        ("Article2", "Second Article"),
        ("Article3", "Third Article"),
    ]):
        d = tmp_path / folder
        d.mkdir()
        _write_v3_doc(d / f"doc{i}.docx", title)

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]

    result = run_batch(files, page_type_config=_sample_config())

    assert len(result.entries) == 3
    folders = [e.folder for e in result.entries]
    assert folders == ["Article1", "Article2", "Article3"]
    assert all(e.status == "passed" for e in result.entries)
    # page path derived from meta title, subfolder name absent
    assert result.entries[0].page_path == "/content/site/articles/first-article"
    assert result.entries[1].page_path == "/content/site/articles/second-article"
    assert "article1" not in result.entries[0].page_path.lower()


def test_batch_entry_carries_full_aem_tree(tmp_path):
    (tmp_path / "Article1").mkdir()
    _write_v3_doc(tmp_path / "Article1" / "a.docx", "The Only Article")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in tmp_path.rglob("*.docx")]

    result = run_batch(files, page_type_config=_sample_config())

    assert result.entries[0].aem_tree is not None
    assert result.entries[0].aem_tree["primary_type"] == "cq:Page"


def test_batch_flags_page_path_collisions(tmp_path):
    """Two docs with identical meta title -> same JCR path. First wins,
    second recorded as a collision -- no silent overwrite."""
    for folder in ["A1", "A2"]:
        d = tmp_path / folder
        d.mkdir()
        _write_v3_doc(d / "d.docx", "Same Title")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]

    result = run_batch(files, page_type_config=_sample_config())

    assert result.entries[0].status == "passed"
    assert result.entries[1].status == "collision"
    assert "same-title" in result.entries[1].error.lower()
    assert result.entries[1].aem_tree is None


def test_batch_survives_a_failing_doc_and_continues(tmp_path):
    (tmp_path / "A1").mkdir()
    _write_v3_doc(tmp_path / "A1" / "ok.docx", "Good One")
    # a broken file (not a real docx)
    (tmp_path / "A2").mkdir()
    (tmp_path / "A2" / "broken.docx").write_bytes(b"not a real docx")
    (tmp_path / "A3").mkdir()
    _write_v3_doc(tmp_path / "A3" / "also-ok.docx", "Good Two")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]

    result = run_batch(files, page_type_config=_sample_config())

    assert [e.status for e in result.entries] == ["passed", "failed", "passed"]
    # failed entry keeps its identity so operator sees which one broke
    assert result.entries[1].folder == "A2"
    assert result.entries[1].filename == "broken.docx"


def _write_v3_doc_no_seed(path: Path) -> None:
    """v3 doc that validates fine (has Meta Title + Meta Description) but
    whose Meta Title slugifies to only non-alphanumeric characters, so
    the mapper's pageNameStrategy = slug_from_meta_title falls back to
    the UNRESOLVED_SLUG placeholder."""
    document = docx.Document()
    document.add_paragraph().add_run("H1").bold = True
    # Meta Title contains only punctuation -> slugify collapses to empty
    document.add_paragraph("Meta Title: ---")
    document.add_paragraph("Meta Description: desc")
    document.add_paragraph("Meta keywords: kw")
    document.add_paragraph("Body.")
    document.save(str(path))


def test_batch_flags_unresolvable_page_names_and_drops_them(tmp_path):
    """A doc that validates OK but whose title slugifies to the
    UNRESOLVED_SLUG placeholder ('page'). Batch mode marks it unresolved,
    drops the aem_tree so it never lands in the package, and keeps the
    row in the log so the operator sees which doc to fix."""
    (tmp_path / "A1").mkdir()
    _write_v3_doc_no_seed(tmp_path / "A1" / "nameless.docx")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in tmp_path.rglob("*.docx")]

    result = run_batch(files, page_type_config=_sample_config())

    assert result.entries[0].status == "unresolved"
    # dropped from the package -- no aem_tree available for the batch builder
    assert result.entries[0].aem_tree is None
    assert "could not derive" in result.entries[0].error.lower()


def test_batch_unresolved_pages_are_excluded_from_slug_collision_set(tmp_path):
    """Three docs all slugging to the placeholder must not eat each other
    as 'collisions' -- each is independently unresolved."""
    for folder in ["A1", "A2", "A3"]:
        d = tmp_path / folder
        d.mkdir()
        _write_v3_doc_no_seed(d / "d.docx")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]

    result = run_batch(files, page_type_config=_sample_config())

    assert [e.status for e in result.entries] == ["unresolved", "unresolved", "unresolved"]


def _slug_from_url_config(site_origin: str = "https://www.iffcotokio.co.in") -> PageTypeConfig:
    return PageTypeConfig(
        type="article",
        template_path="/conf/x",
        content_root="/content/site/articles",
        dam_root="/content/dam/site/articles",
        page_name_strategy="slug_from_url",
        component_mappings={"page-description": "site/components/text"},
        site_origin=site_origin,
    )


def _write_v3_doc_with_url(path: Path, title: str, url: str | None) -> None:
    document = docx.Document()
    document.add_paragraph().add_run(title).bold = True
    if url:
        document.add_paragraph(f"New Required URL: {url}")
    document.add_paragraph(f"Meta Title: {title}")
    document.add_paragraph("Meta Description: Some description.")
    document.add_paragraph("Meta keywords: kw")
    document.add_paragraph("Body.")
    document.save(str(path))


def test_batch_unresolved_reason_names_missing_url_under_slug_from_url(tmp_path):
    """slug_from_url with no URL at all in the doc -- the message must say
    the doc has no URL, not point at an unrelated config value."""
    (tmp_path / "A1").mkdir()
    _write_v3_doc_with_url(tmp_path / "A1" / "no-url.docx", "Some Title", url=None)

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in tmp_path.rglob("*.docx")]

    result = run_batch(files, page_type_config=_slug_from_url_config())

    assert result.entries[0].status == "unresolved"
    reason = result.entries[0].error.lower()
    assert "no page url" in reason
    assert "siteorigin" not in reason


def test_batch_unresolved_reason_names_site_origin_mismatch_under_slug_from_url(tmp_path):
    """slug_from_url with a URL present but not on the configured
    siteOrigin -- the message must name siteOrigin as the thing to check,
    not blame the doc's own metadata."""
    (tmp_path / "A1").mkdir()
    _write_v3_doc_with_url(
        tmp_path / "A1" / "wrong-origin.docx", "Some Title",
        url="https://example.com/health-insurance/who-buys")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in tmp_path.rglob("*.docx")]

    result = run_batch(files, page_type_config=_slug_from_url_config(
        site_origin="https://www.iffcotokio.co.in"))

    assert result.entries[0].status == "unresolved"
    reason = result.entries[0].error.lower()
    assert "siteorigin" in reason
    assert "https://example.com/health-insurance/who-buys" in result.entries[0].error


def test_batch_result_carries_top_folder_label(tmp_path):
    (tmp_path / "Articles" / "A1").mkdir(parents=True)
    _write_v3_doc(tmp_path / "Articles" / "A1" / "doc.docx", "One")

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in tmp_path.rglob("*.docx")]

    result = run_batch(files, page_type_config=_sample_config())
    # top folder = first segment before subfolder + filename
    assert result.filename_top == "Articles"


# ---- Phase 2B: folder assets threaded through the pipeline ----

def _hero_banner_cfg() -> PageTypeConfig:
    """Aligned with real client shape: template overlay with a hero
    banner slot, mapping bound to `new-hero-banner`."""
    return PageTypeConfig(
        type="article",
        template_path="/conf/x",
        content_root="/content/site/articles",
        dam_root="/content/dam/site/articles",
        page_name_strategy="slug_from_meta_title",
        component_mappings={
            "new-hero-banner": {
                "resourceType": "site/components/hero",
                "fieldRenames": {"htmltext": "bannerDescription"},
                "staticProperties": {
                    "cardAlign": "left",
                    "imagePath": "/content/dam/site/banner/default.png",
                },
            },
        },
        component_slots={
            "new-hero-banner": "root/responsivegrid/new_hero_banner",
        },
        template_initial_content={
            "name": "jcr:content", "primary_type": "cq:PageContent",
            "properties": {"sling:resourceType": "site/components/page"},
            "children": [{
                "name": "root", "primary_type": "nt:unstructured",
                "properties": {"sling:resourceType": "wcm/foundation/components/responsivegrid"},
                "children": [{
                    "name": "responsivegrid", "primary_type": "nt:unstructured",
                    "properties": {"sling:resourceType": "wcm/foundation/components/responsivegrid"},
                    "children": [{
                        "name": "new_hero_banner",
                        "primary_type": "nt:unstructured",
                        "properties": {
                            "sling:resourceType": "site/components/hero",
                            "bannerTitle": "Template Default",
                        },
                        "children": [],
                    }],
                }],
            }],
        },
    )


def _write_hero_doc(path: Path, title: str, image_ref: str) -> None:
    """Minimal v1/v2 doc that validates: `[metadata]` table (title +
    description), then a `[component:new-hero-banner]` block with the
    Banner Image marker and a description paragraph."""
    document = docx.Document()
    document.add_paragraph("[metadata]")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Property", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "title", title
    table.cell(2, 0).text, table.cell(2, 1).text = "description", "Some desc"
    p = document.add_paragraph()
    p.add_run("[component:new-hero-banner]").add_break()
    p.add_run(f"Banner Image: {image_ref}").add_break()
    p.add_run().add_break()
    p.add_run("A short banner description.")
    document.save(str(path))


def test_batch_wires_folder_assets_to_hero_banner_imagepath(tmp_path):
    """Non-docx files in each subfolder ride the batch pipeline as page
    assets. A hero-banner marker referencing one of those files -> that
    file resolves to a DAM path on the mapped page, and an asset
    descriptor is attached to the aem_tree for the packager."""
    (tmp_path / "Article1").mkdir()
    _write_hero_doc(tmp_path / "Article1" / "V2.docx",
                    "Cashless Treatment", "car-insurance.jpg")
    (tmp_path / "Article1" / "car-insurance.jpg").write_bytes(b"\x89PNG-bytes")

    files = [
        (str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
        for f in sorted(tmp_path.rglob("*")) if f.is_file()
    ]

    result = run_batch(files, page_type_config=_hero_banner_cfg())

    # only one docx -> one pipeline entry (the jpg is an asset, not a page)
    assert len(result.entries) == 1
    entry = result.entries[0]
    assert entry.status == "passed"

    # hero banner has imagePath set to the DAM path for the folder image
    root = entry.aem_tree["children"][0]["children"][0]
    grid = next(c for c in root["children"] if c["name"] == "responsivegrid")
    hero = next(c for c in grid["children"] if c["name"] == "new_hero_banner")
    assert hero["properties"]["imagePath"] == \
        "/content/dam/site/articles/car-insurance.jpg"

    # asset descriptor rides on the entry (bytes -> base64 at DB persist);
    # the batch package route re-attaches it to a tree copy on download.
    assert len(entry.assets) == 1
    assert entry.assets[0]["filename"] == "car-insurance.jpg"
    assert entry.assets[0]["bytes"] == b"\x89PNG-bytes"
    # ...and MUST have been stripped off the tree itself (not JSON-safe).
    assert "_assets" not in entry.aem_tree


def test_batch_missing_folder_asset_keeps_template_default(tmp_path):
    """Doc references `Banner Image: gone.jpg` but the folder has a
    different image. Pipeline still produces a page (no failure), the
    template default imagePath survives, no asset descriptor attached."""
    (tmp_path / "Article1").mkdir()
    _write_hero_doc(tmp_path / "Article1" / "V2.docx",
                    "Article Title", "gone.jpg")
    (tmp_path / "Article1" / "other.jpg").write_bytes(b"unrelated")

    files = [
        (str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
        for f in sorted(tmp_path.rglob("*")) if f.is_file()
    ]

    result = run_batch(files, page_type_config=_hero_banner_cfg())

    entry = result.entries[0]
    assert entry.status == "passed"
    root = entry.aem_tree["children"][0]["children"][0]
    grid = next(c for c in root["children"] if c["name"] == "responsivegrid")
    hero = next(c for c in grid["children"] if c["name"] == "new_hero_banner")
    # template default (from staticProperties) preserved
    assert hero["properties"]["imagePath"] == \
        "/content/dam/site/banner/default.png"
    assert entry.assets == []


def test_batch_assets_scoped_per_folder_no_cross_leak(tmp_path):
    """Article1's image must not resolve for Article2's banner marker.
    Assets are grouped by first path segment (the subfolder name)."""
    (tmp_path / "Article1").mkdir()
    _write_hero_doc(tmp_path / "Article1" / "v.docx", "A1 Title", "img.jpg")
    (tmp_path / "Article1" / "img.jpg").write_bytes(b"a1-bytes")

    (tmp_path / "Article2").mkdir()
    _write_hero_doc(tmp_path / "Article2" / "v.docx", "A2 Title", "img.jpg")
    # no image in Article2 -- should NOT pick up Article1's img.jpg

    files = [
        (str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
        for f in sorted(tmp_path.rglob("*")) if f.is_file()
    ]

    result = run_batch(files, page_type_config=_hero_banner_cfg())

    a1_entry = next(e for e in result.entries if e.folder == "Article1")
    a2_entry = next(e for e in result.entries if e.folder == "Article2")

    a1_hero_path = _dig_hero_imagepath(a1_entry.aem_tree)
    a2_hero_path = _dig_hero_imagepath(a2_entry.aem_tree)

    assert a1_hero_path == "/content/dam/site/articles/img.jpg"
    # Article2's marker doesn't resolve -> template default survives
    assert a2_hero_path == "/content/dam/site/banner/default.png"
    assert a2_entry.assets == []


def _dig_hero_imagepath(tree: dict) -> str:
    root = tree["children"][0]["children"][0]
    grid = next(c for c in root["children"] if c["name"] == "responsivegrid")
    hero = next(c for c in grid["children"] if c["name"] == "new_hero_banner")
    return hero["properties"].get("imagePath", "")


# ---- Task 7: FAQ Content Fragments carried through batch entries ----

def test_batch_pops_faq_cfs_but_leaves_cf_root_and_source_url_on_tree(tmp_path, monkeypatch):
    """`_faq_cfs` is JSON-safe (Task 6) but still rides on `BatchEntry.faq_cfs`,
    same as `_assets` -> so `entries_to_json_list` round-trips it without
    a base64 step. `_faq_cf_root` / `_source_url` are plain strings and
    must survive untouched on `aem_tree` (the packager needs the former,
    Task 9 reads the latter)."""
    from ui.service import RunResult

    d = tmp_path / "Article1"
    d.mkdir()
    doc_path = d / "doc.docx"
    _write_v3_doc(doc_path, "FAQ Page")

    fake_tree = {
        "name": "p", "path": "/content/site/articles/faq-page",
        "primary_type": "cq:Page", "properties": {}, "children": [],
        "_faq_cf_root": "/content/dam/site/faq",
        "_source_url": "https://example.com/faq-page",
        "_faq_cfs": [{"name": "q-one", "path": "/content/dam/site/faq/q-one",
                      "primary_type": "dam:Asset", "properties": {}, "children": []}],
    }
    monkeypatch.setattr(
        "ui.batch_service.run_pipeline",
        lambda *a, **kw: RunResult(status="passed", filename="doc.docx",
                                    issues=[], ncm={}, aem_tree=fake_tree),
    )

    result = run_batch([(f"Article1/{doc_path.name}", doc_path.read_bytes())],
                        page_type_config=_sample_config())

    entry = result.entries[0]
    assert entry.faq_cfs == [{"name": "q-one", "path": "/content/dam/site/faq/q-one",
                               "primary_type": "dam:Asset", "properties": {}, "children": [],
                               "collision": "unknown"}]
    assert "_faq_cfs" not in entry.aem_tree
    assert entry.aem_tree["_faq_cf_root"] == "/content/dam/site/faq"
    assert entry.aem_tree["_source_url"] == "https://example.com/faq-page"


def test_batch_entries_json_round_trip_preserves_faq_cfs(tmp_path, monkeypatch):
    from ui.batch_service import entries_from_json_list, entries_to_json_list
    from ui.service import RunResult

    d = tmp_path / "Article1"
    d.mkdir()
    doc_path = d / "doc.docx"
    _write_v3_doc(doc_path, "FAQ Page")

    fake_tree = {
        "name": "p", "path": "/content/site/articles/faq-page",
        "primary_type": "cq:Page", "properties": {}, "children": [],
        "_faq_cf_root": "/content/dam/site/faq",
        "_faq_cfs": [{"name": "q-one", "path": "/content/dam/site/faq/q-one",
                      "primary_type": "dam:Asset", "properties": {}, "children": []}],
    }
    monkeypatch.setattr(
        "ui.batch_service.run_pipeline",
        lambda *a, **kw: RunResult(status="passed", filename="doc.docx",
                                    issues=[], ncm={}, aem_tree=fake_tree),
    )

    result = run_batch([(f"Article1/{doc_path.name}", doc_path.read_bytes())],
                        page_type_config=_sample_config())

    rows = entries_to_json_list(result.entries)
    restored = entries_from_json_list(rows)

    assert restored[0].faq_cfs == result.entries[0].faq_cfs


# ---- Task 8: collision detection against live AEM ----

def _fake_tree_with_one_faq_cf():
    return {
        "name": "p", "path": "/content/site/articles/faq-page",
        "primary_type": "cq:Page", "properties": {}, "children": [],
        "_faq_cf_root": "/content/dam/site/faq",
        "_faq_cfs": [{"name": "q-one", "path": "/content/dam/site/faq/q-one",
                      "primary_type": "dam:Asset", "properties": {}, "children": []}],
    }


def test_batch_with_no_aem_base_url_stamps_unknown_and_makes_no_network_call(
    tmp_path, monkeypatch,
):
    """No AEM configured -> every fragment reports "unknown" honestly,
    and _annotate_collisions must not touch the network to get there."""
    from ui.service import RunResult

    d = tmp_path / "Article1"
    d.mkdir()
    doc_path = d / "doc.docx"
    _write_v3_doc(doc_path, "FAQ Page")

    monkeypatch.setattr(
        "ui.batch_service.run_pipeline",
        lambda *a, **kw: RunResult(status="passed", filename="doc.docx",
                                    issues=[], ncm={}, aem_tree=_fake_tree_with_one_faq_cf()),
    )

    def _no_network(*args, **kwargs):
        raise AssertionError("cf_exists must not be called with no aem_base_url")
    monkeypatch.setattr("ui.batch_service.cf_exists", _no_network)

    result = run_batch([(f"Article1/{doc_path.name}", doc_path.read_bytes())],
                        page_type_config=_sample_config())

    assert result.entries[0].faq_cfs[0]["collision"] == "unknown"


def test_batch_with_aem_base_url_annotates_each_faq_cf_via_cf_exists(
    tmp_path, monkeypatch,
):
    from ui.service import RunResult

    d = tmp_path / "Article1"
    d.mkdir()
    doc_path = d / "doc.docx"
    _write_v3_doc(doc_path, "FAQ Page")

    monkeypatch.setattr(
        "ui.batch_service.run_pipeline",
        lambda *a, **kw: RunResult(status="passed", filename="doc.docx",
                                    issues=[], ncm={}, aem_tree=_fake_tree_with_one_faq_cf()),
    )

    calls = []

    def _fake_cf_exists(client, base_url, cf_path, auth_token=None):
        calls.append((base_url, cf_path, auth_token))
        return "yes"
    monkeypatch.setattr("ui.batch_service.cf_exists", _fake_cf_exists)

    result = run_batch(
        [(f"Article1/{doc_path.name}", doc_path.read_bytes())],
        page_type_config=_sample_config(),
        aem_base_url="http://localhost:4502",
        auth_token="tok",
    )

    assert result.entries[0].faq_cfs[0]["collision"] == "yes"
    assert calls == [("http://localhost:4502", "/content/dam/site/faq/q-one", "tok")]


# ---- Final review, Finding 1: cross-document CF path collision ----

def _faq_config() -> PageTypeConfig:
    cfg = _sample_config()
    cfg.page_name_strategy = "slug_from_url"
    cfg.content_root = "/content/site/en"
    cfg.site_origin = "https://www.example.com"
    cfg.faq = FaqConfig(cf_root="/content/dam/site/faq", cf_model="/conf/m",
                        tag_namespace="site")
    return cfg


def _write_faq_doc(path: Path, title: str, url: str, qa: list[tuple[str, str]]) -> None:
    document = docx.Document()
    document.add_paragraph().add_run(title).bold = True
    document.add_paragraph(f"Meta Title: {title}")
    document.add_paragraph("Meta Description: Some description.")
    document.add_paragraph(f"New Required URL: {url}")
    document.add_paragraph("Body content.")
    document.add_paragraph().add_run("FAQs").bold = True
    for question, answer in qa:
        document.add_paragraph().add_run(question).bold = True
        document.add_paragraph(answer)
    document.save(str(path))


def test_two_docs_sharing_a_question_keep_both_answers_and_report_the_clash(tmp_path):
    """Two documents in the same URL section asking the same question derive
    the SAME fragment path. The packager keys entries by path, so without
    batch-wide uniqueness the second answer overwrites the first and one
    answer ships missing with no issue, no error and no flag. Both answers
    must survive AND the clash must be visible to the operator."""
    import zipfile
    from io import BytesIO

    from aem.package_generator.package import PackageMeta, build_batch_package
    from ui.report import build_report_zip

    question = "What is a family floater plan?"
    for folder, name, slug, answer in [
        ("Article1", "One", "one", "One policy covers the whole family."),
        ("Article2", "Two", "two", "A totally different explanation entirely."),
    ]:
        d = tmp_path / folder
        d.mkdir()
        _write_faq_doc(d / f"{slug}.docx", name,
                       f"https://www.example.com/health-insurance/{slug}",
                       [(question, answer)])

    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]

    result = run_batch(files, page_type_config=_faq_config())

    assert [e.status for e in result.entries] == ["passed", "passed"]
    paths = [cf["path"] for e in result.entries for cf in e.faq_cfs]
    assert len(paths) == 2
    assert len(set(paths)) == 2, "fragment paths must be unique batch-wide"

    # 1. no answer is lost from the package
    trees = []
    for entry in result.entries:
        tree = {**entry.aem_tree, "_faq_cfs": entry.faq_cfs}
        trees.append(tree)
    data = build_batch_package(trees, PackageMeta(name="batch"))
    with zipfile.ZipFile(BytesIO(data)) as zf:
        xml = "".join(zf.read(n).decode() for n in zf.namelist()
                      if n.endswith(".content.xml"))
        filter_xml = zf.read("META-INF/vault/filter.xml").decode()
    assert "One policy covers the whole family." in xml
    assert "A totally different explanation entirely." in xml
    # one filter root per fragment, none repeated
    roots = [line for line in filter_xml.splitlines() if "<filter root=" in line]
    assert len(roots) == len(set(roots)) == 4      # 2 pages + 2 fragments

    # 2. the clash is visible in the report
    import csv as csv_mod
    import io as io_mod
    with zipfile.ZipFile(BytesIO(build_report_zip(result))) as zf:
        faqs = list(csv_mod.DictReader(io_mod.StringIO(
            zf.read("faqs.csv").decode("utf-8-sig"))))
        pages = list(csv_mod.DictReader(io_mod.StringIO(
            zf.read("pages.csv").decode("utf-8-sig"))))
    clashed = [r for r in faqs if r["batch_clash"]]
    assert len(clashed) == 1
    assert clashed[0]["batch_clash"] == \
        "/content/dam/site/faq/health-insurance/what-is-a-family-floater-plan"
    assert clashed[0]["cf_path"] == clashed[0]["batch_clash"] + "-2"
    assert any("clashed" in row["warnings"] for row in pages)


def test_faq_reference_on_the_page_follows_a_renamed_fragment(tmp_path):
    """The page's fragmentPaths property is built from the pre-rename paths --
    it has to be rewritten too, or the page points at a fragment this batch
    never writes."""
    cfg = _faq_config()
    cfg.faq.page_component = {"resourceType": "site/components/faq",
                              "pathsProperty": "fragmentPaths"}
    cfg.component_mappings = {**cfg.component_mappings, "faq": "site/components/faq"}
    question = "What is a family floater plan?"
    for folder, slug in [("Article1", "one"), ("Article2", "two")]:
        d = tmp_path / folder
        d.mkdir()
        _write_faq_doc(d / f"{slug}.docx", slug.title(),
                       f"https://www.example.com/health-insurance/{slug}",
                       [(question, f"Answer for {slug}.")])
    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]

    result = run_batch(files, page_type_config=cfg)

    second = result.entries[1]
    referenced = []

    def walk(node):
        for key, value in (node.get("properties") or {}).items():
            if key == "fragmentPaths":
                referenced.extend(value)
        for child in node.get("children") or []:
            walk(child)
    walk(second.aem_tree)

    assert referenced == [cf["path"] for cf in second.faq_cfs]
    assert referenced[0].endswith("-2")


def _faq_ref_config() -> PageTypeConfig:
    cfg = _faq_config()
    cfg.faq.page_component = {"resourceType": "site/components/faq",
                              "pathsProperty": "fragmentPaths"}
    cfg.component_mappings = {**cfg.component_mappings, "faq": "site/components/faq"}
    return cfg


def _fragment_paths(tree: dict) -> list[str]:
    found: list[str] = []

    def walk(node):
        for key, value in (node.get("properties") or {}).items():
            if key == "fragmentPaths":
                found.extend(value)
        for child in node.get("children") or []:
            walk(child)
    walk(tree)
    return found


def _shipped_cf_paths(entries) -> set[str]:
    """CF paths the package actually writes, read back out of the zip."""
    import zipfile
    from io import BytesIO

    from aem.package_generator.package import PackageMeta, build_batch_package

    trees = [{**e.aem_tree, "_faq_cfs": e.faq_cfs} for e in entries]
    with zipfile.ZipFile(BytesIO(build_batch_package(trees, PackageMeta(name="b")))) as zf:
        names = zf.namelist()
    return {"/" + n[len("jcr_root/"):-len("/.content.xml")]
            for n in names
            if n.startswith("jcr_root/content/dam/") and n.endswith("/.content.xml")}


def _run_two_docs(tmp_path, doc_a_qa, doc_b_qa):
    for folder, slug, qa in [("Article1", "one", doc_a_qa), ("Article2", "two", doc_b_qa)]:
        d = tmp_path / folder
        d.mkdir()
        _write_faq_doc(d / f"{slug}.docx", slug.title(),
                       f"https://www.example.com/health-insurance/{slug}", qa)
    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in sorted(tmp_path.rglob("*.docx"))]
    return run_batch(files, page_type_config=_faq_ref_config())


def test_page_refs_survive_two_renames_in_one_document(tmp_path):
    """Doc B asks the clashing question twice, so BOTH its fragments are
    renamed. Rewriting the page property one old->new value at a time made
    the second rewrite re-match what the first had just written: one fragment
    referenced twice, the other shipped and orphaned. The page's reference
    list must name exactly the fragments the package writes for that page."""
    question = "What is a family floater plan?"
    result = _run_two_docs(
        tmp_path,
        [(question, "Doc A answer.")],
        [(question, "Doc B answer."), (question, "Doc B second answer.")])

    assert [e.status for e in result.entries] == ["passed", "passed"]
    shipped = _shipped_cf_paths(result.entries)
    second = result.entries[1]
    refs = _fragment_paths(second.aem_tree)

    assert refs == [cf["path"] for cf in second.faq_cfs]
    assert len(refs) == len(set(refs)) == 2          # no duplicate reference
    assert set(refs) <= shipped                      # no dangling reference
    # every fragment this page ships is referenced by it -- nothing orphaned
    assert {cf["path"] for cf in second.faq_cfs} <= set(refs)


def test_page_refs_survive_a_rename_that_collides_with_a_real_sibling_slug(tmp_path):
    """No in-document duplicate needed, which makes this the easy one to hit
    by accident: doc B's second question genuinely slugs to `...-2`, the same
    name the rename of its first question wants. Value substitution turned
    both references into the same path."""
    result = _run_two_docs(
        tmp_path,
        [("Is it covered?", "Doc A answer.")],
        [("Is it covered?", "Doc B answer."),
         ("Is it covered 2?", "Doc B other answer.")])

    shipped = _shipped_cf_paths(result.entries)
    second = result.entries[1]
    refs = _fragment_paths(second.aem_tree)

    assert [cf["path"] for cf in second.faq_cfs] == [
        "/content/dam/site/faq/health-insurance/is-it-covered-2",
        "/content/dam/site/faq/health-insurance/is-it-covered-2-2",
    ]
    assert refs == [cf["path"] for cf in second.faq_cfs]
    assert len(refs) == len(set(refs)) == 2
    assert set(refs) <= shipped
    assert {cf["path"] for cf in second.faq_cfs} <= set(refs)


# ---- collision key: full page path, not the last segment ----

def _url_docs(tmp_path, specs: list[tuple[str, str, str]]):
    """specs is (folder, title, url). Returns the run_batch files list."""
    for folder, title, url in specs:
        (tmp_path / folder).mkdir(parents=True, exist_ok=True)
        _write_v3_doc_with_url(tmp_path / folder / "d.docx", title, url)
    return [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
            for f in sorted(tmp_path.rglob("*.docx"))]


def test_batch_same_last_segment_under_different_hierarchies_is_not_a_collision(tmp_path):
    """Two pages whose URLs differ only by a language prefix resolve to
    genuinely different JCR paths. Keying the duplicate check on the last
    path segment alone flagged the second as a collision and dropped it
    from the package -- which silently removed every Hindi page from a
    batch that also carried its English counterpart."""
    files = _url_docs(tmp_path, [
        ("en", "Cashless Claim",
         "https://www.iffcotokio.co.in/health-insurance/cashless-treatment"),
        ("hi", "Cashless Claim Hindi",
         "https://www.iffcotokio.co.in/hindi/health-insurance/cashless-treatment"),
    ])

    result = run_batch(files, page_type_config=_slug_from_url_config())

    assert [e.status for e in result.entries] == ["passed", "passed"]
    assert sorted(e.page_path for e in result.entries) == [
        "/content/site/articles/health-insurance/cashless-treatment",
        "/content/site/articles/hindi/health-insurance/cashless-treatment",
    ]
    # both survive into the package
    assert all(e.aem_tree is not None for e in result.entries)


def test_batch_identical_page_paths_still_collide(tmp_path):
    """The duplicate guard must still fire when two documents resolve to
    the SAME full path -- widening the key must not disable the check."""
    files = _url_docs(tmp_path, [
        ("a", "Cashless Claim",
         "https://www.iffcotokio.co.in/health-insurance/cashless-treatment"),
        ("b", "Cashless Claim Again",
         "https://www.iffcotokio.co.in/health-insurance/cashless-treatment"),
    ])

    result = run_batch(files, page_type_config=_slug_from_url_config())

    assert [e.status for e in result.entries] == ["passed", "collision"]
    assert result.entries[1].aem_tree is None
    # the message names the full path, so the operator can find both docs
    assert "/content/site/articles/health-insurance/cashless-treatment" \
        in result.entries[1].error


def test_batch_without_a_page_type_config_reports_unmapped_not_passed(tmp_path):
    """With no resolvable PageTypeConfig the pipeline produces no node
    tree, so the page cannot reach the package. It used to be recorded as
    `passed` with a blank path and no error, so the operator saw a green
    batch and silently got an empty package. It must say what happened."""
    (tmp_path / "A1").mkdir()
    _write_v3_doc(tmp_path / "A1" / "d.docx", "Some Page")
    files = [(str(f.relative_to(tmp_path)).replace("\\", "/"), f.read_bytes())
             for f in tmp_path.rglob("*.docx")]

    result = run_batch(files, page_type_config=None)

    entry = result.entries[0]
    assert entry.status == "unmapped"
    assert entry.aem_tree is None
    assert "config" in (entry.error or "").lower()
    # the row still names the file so the operator can see what was read
    assert entry.filename == "d.docx"
